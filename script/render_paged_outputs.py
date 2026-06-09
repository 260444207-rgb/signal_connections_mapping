#!/usr/bin/env python3
# -*- coding: utf-8 -*-
from __future__ import annotations
import argparse,csv,re
from pathlib import Path
from typing import Dict, Any, List
from common import iter_jsonl,write_jsonl,ensure_dir,FINAL_HEADERS

def safe_sheet_name(name: str) -> str:
    name = str(name or 'UNMAPPED').strip()
    name = re.sub(r'[\\/*?:\[\]]+', '_', name)
    return name[:31] or 'UNMAPPED'

def build_final_rows(normalized_path,decisions_path):
    normalized={r['line_id']:r for r in iter_jsonl(normalized_path)}
    decisions={d['line_id']:d for d in iter_jsonl(decisions_path)}
    rows=[]
    for line,n in normalized.items():
        d=decisions.get(line,{})
        rows.append({
            'output_sheet_name': n.get('output_sheet_name') or n.get('source_sheet_name') or n.get('source_block_id',''),
            'source_sheet_name': n.get('source_sheet_name',''),
            'line_id': line,
            '起止Block标识':n.get('source_block_id',''),
            '起始Block名称':n.get('source_block_name',''),
            '源Port':n.get('source_port',''),
            '目标Block标识':n.get('target_block_id',''),
            '目标Block名称':n.get('target_block_name',''),
            '目标Port':n.get('target_port',''),
            '连线标识':n.get('connection_id',''),
            '连线名称':n.get('connection_name',''),
            '连线方向':n.get('direction',''),
            '源原理图Pin脚':d.get('selected_pin',''),
            '分析说明':d.get('analysis',''),
            '映射置信度':d.get('confidence',''),
            '网络命名':d.get('net_name','')
        })
    return rows

def group_rows_by_sheet(rows: List[Dict[str,Any]]) -> Dict[str,List[Dict[str,Any]]]:
    grouped: Dict[str,List[Dict[str,Any]]] = {}
    for row in rows:
        sheet = row.get('output_sheet_name') or row.get('source_sheet_name') or row.get('起止Block标识') or 'UNMAPPED'
        grouped.setdefault(str(sheet),[]).append(row)
    return grouped

def clear_sheet(ws):
    if ws.max_row:
        ws.delete_rows(1, ws.max_row)

def write_sheet(ws, rows):
    from openpyxl.styles import Font,Alignment,PatternFill
    header_fill=PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    ws.append(FINAL_HEADERS)
    for c in ws[1]:
        c.font=Font(bold=True,color='FFFFFF')
        c.fill=header_fill
        c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
    for row in rows:
        ws.append([row.get(h,'') for h in FINAL_HEADERS])
    for col in ws.columns:
        ml=max(len(str(cell.value or '')) for cell in col)
        ws.column_dimensions[col[0].column_letter].width=min(42,max(10,ml+2))

def render_paged_xlsx(rows, output_path, template_excel=''):
    from openpyxl import Workbook, load_workbook
    p=Path(output_path); ensure_dir(p.parent)
    grouped=group_rows_by_sheet(rows)
    if template_excel:
        wb=load_workbook(template_excel)
    else:
        wb=Workbook()
        default=wb.active
        wb.remove(default)
    for raw_sheet, sheet_rows in grouped.items():
        sheet_name=safe_sheet_name(raw_sheet)
        if sheet_name in wb.sheetnames:
            ws=wb[sheet_name]
            clear_sheet(ws)
        else:
            ws=wb.create_sheet(sheet_name)
        write_sheet(ws, sheet_rows)
    if not wb.sheetnames:
        ws=wb.create_sheet('EMPTY')
        write_sheet(ws, [])
    wb.save(p)
    wb.close()

def render_per_sheet_markdown(rows, output_dir):
    out=Path(output_dir)/'mapping_by_device'; ensure_dir(out)
    for sheet, sheet_rows in group_rows_by_sheet(rows).items():
        path=out/f'{safe_sheet_name(sheet)}_mapping.md'
        with path.open('w',encoding='utf-8') as f:
            f.write('# 映射结果\n\n')
            f.write('| '+' | '.join(FINAL_HEADERS)+' |\n')
            f.write('|'+'|'.join(['---']*len(FINAL_HEADERS))+'|\n')
            for row in sheet_rows:
                f.write('| '+' | '.join([str(row.get(h,'') or '').replace('|','/') for h in FINAL_HEADERS])+' |\n')

def render_csv(rows, output_path):
    p=Path(output_path); ensure_dir(p.parent)
    headers=['output_sheet_name','source_sheet_name','line_id']+FINAL_HEADERS
    with p.open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.DictWriter(f,fieldnames=headers); w.writeheader(); w.writerows(rows)

def render_paged_outputs(normalized_path, decisions_path, output_dir, template_excel=''):
    out=Path(output_dir); ensure_dir(out)
    rows=build_final_rows(normalized_path, decisions_path)
    write_jsonl(out/'final_mapping_rows.jsonl', rows)
    render_csv(rows, out/'mapping_result_flat_debug.csv')
    render_per_sheet_markdown(rows, out)
    render_paged_xlsx(rows, out/'signal_interface_paged.xlsx', template_excel)

def main():
    p=argparse.ArgumentParser()
    p.add_argument('--normalized',required=True)
    p.add_argument('--decisions',required=True)
    p.add_argument('--output-dir',required=True)
    p.add_argument('--template-excel',default='')
    a=p.parse_args()
    render_paged_outputs(a.normalized,a.decisions,a.output_dir,a.template_excel)
    print(f'[OK] rendered paged signal interface -> {a.output_dir}')
if __name__=='__main__': main()
