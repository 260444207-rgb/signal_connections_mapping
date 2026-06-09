#!/usr/bin/env python3
# -*- coding: utf-8 -*-

from __future__ import annotations

import argparse
import re
from pathlib import Path
from typing import Dict, List, Any, Iterable

from common import write_jsonl, normalize_text, normalize_direction, read_table

FIELD_ALIASES = {
    "source_block_id": ["source_block_id", "源Block标识", "起止Block标识", "起始Block标识", "source_id", "block_id", "框图标识"],
    "source_block_name": ["source_block_name", "源Block名称", "起始Block名称", "source_name", "器件名称"],
    "source_port": ["source_port", "源Port", "源端Port", "port"],
    "target_block_id": ["target_block_id", "目的Block标识", "目标Block标识", "target_id"],
    "target_block_name": ["target_block_name", "目的Block名称", "目标Block名称", "target_name"],
    "target_port": ["target_port", "目的Port", "目标Port"],
    "connection_id": ["connection_id", "连线ID", "连线标识", "line_id", "edge_id"],
    "connection_name": ["connection_name", "连线名称", "net_name"],
    "direction": ["direction", "连线方向"],
}

SKIP_SHEETS = {"BLOCK_INFO", "说明", "README", "INDEX", "目录"}

def pick(row: Dict[str, Any], canonical: str) -> str:
    for k in FIELD_ALIASES[canonical]:
        if k in row and row[k] is not None:
            return normalize_text(row[k])
    return ""

def expansion_count(port: str) -> int:
    port = normalize_text(port)
    m = re.match(r"^(.+)\[(\d+):(\d+)\]$", port)
    if m:
        a, b = int(m.group(2)), int(m.group(3))
        return abs(a - b) + 1
    m = re.match(r"^(.+)\*(\d+)$", port)
    if m:
        return int(m.group(2))
    return 1

def read_block_info(path: str | Path) -> Dict[str, str]:
    p = Path(path)
    if p.suffix.lower() not in {".xlsx", ".xlsm"}:
        return {}
    from openpyxl import load_workbook
    wb = load_workbook(p, data_only=True)
    try:
        ws = None
        for candidate in wb.worksheets:
            if candidate.title.strip().lower() == "block_info":
                ws = candidate
                break
        if ws is None:
            return {}
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return {}
        headers = [normalize_text(v) for v in rows[0]]
        try:
            key_idx = headers.index("框图标识")
            part_idx = headers.index("器件信息")
        except ValueError:
            return {}
        result = {}
        for raw in rows[1:]:
            key = normalize_text(raw[key_idx] if key_idx < len(raw) else "")
            part = normalize_text(raw[part_idx] if part_idx < len(raw) else "")
            if key and part:
                result[key.lower()] = part
        return result
    finally:
        wb.close()

def read_excel_sheets(path: str | Path) -> Iterable[tuple[str, list[dict[str, Any]]]]:
    """
    读取输入框图 Excel 的原始 sheet。
    关键原则：输出分页只以这里读取到的 sheet_name 为准，不根据 block_id/group 自动造 sheet。
    """
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    for sheet_name in wb.sheetnames:
        if sheet_name.upper() in SKIP_SHEETS:
            continue
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue

        # 找第一行非空作为表头
        header_idx = None
        headers = None
        for i, r in enumerate(rows):
            values = [normalize_text(v) for v in r]
            if any(values):
                header_idx = i
                headers = values
                break
        if header_idx is None or not headers:
            continue

        data_rows = []
        for r in rows[header_idx + 1:]:
            values = [normalize_text(v) for v in r]
            if not any(values):
                continue
            row = {}
            for h, v in zip(headers, values):
                if h:
                    row[h] = v
            data_rows.append(row)

        yield sheet_name, data_rows

def read_connections_with_sheet(input_path: str | Path) -> Iterable[tuple[str, list[dict[str, Any]]]]:
    p = Path(input_path)
    if p.suffix.lower() in {".xlsx", ".xlsm"}:
        yield from read_excel_sheets(p)
    else:
        # JSON/CSV 也必须显式给 source_sheet_name/output_sheet_name；没有则只作为 DEFAULT_INPUT。
        rows = read_table(p)
        grouped: Dict[str, list[dict[str, Any]]] = {}
        for row in rows:
            sheet = normalize_text(row.get("source_sheet_name") or row.get("output_sheet_name") or row.get("sheet_name") or "DEFAULT_INPUT")
            grouped.setdefault(sheet, []).append(row)
        for sheet, sheet_rows in grouped.items():
            yield sheet, sheet_rows

def normalize_connections(input_path: str | Path, output_path: str | Path) -> List[Dict[str, Any]]:
    out = []
    sheet_parts = read_block_info(input_path)

    for sheet_name, rows in read_connections_with_sheet(input_path):
        for idx, row in enumerate(rows, start=1):
            source_port = pick(row, "source_port")
            if not source_port:
                # 当前 sheet 中非连接行跳过，但不影响最终输出保留该 sheet。
                continue

            target_port_raw = pick(row, "target_port")
            source_block_id = pick(row, "source_block_id") or sheet_name
            source_block_name = pick(row, "source_block_name")
            count = expansion_count(source_port)

            for bit_idx in range(1, count + 1):
                connection_id_raw = pick(row, "connection_id") or f"{sheet_name}_{idx}"
                display_connection_id = connection_id_raw if count == 1 else f"{connection_id_raw}#{bit_idx}"

                normalized = {
                    "line_id": f"{sheet_name}:{idx}:{display_connection_id}",
                    "base_line_id": f"{sheet_name}:{idx}:{connection_id_raw}",
                    "expansion_index": bit_idx,
                    "expansion_count": count,
                    "analysis_unit": "device",
                    "source_sheet_name": sheet_name,
                    "output_sheet_name": sheet_name,
                    "source_part_id": sheet_parts.get(sheet_name.lower(), ""),
                    "source_block_id": source_block_id,
                    "source_block_name": source_block_name,
                    "source_port": source_port,
                    "target_block_id": pick(row, "target_block_id"),
                    "target_block_name": pick(row, "target_block_name"),
                    "target_port": target_port_raw,
                    "connection_id": display_connection_id,
                    "base_connection_id": connection_id_raw,
                    "connection_name": pick(row, "connection_name"),
                    "direction": normalize_direction(pick(row, "direction")),
                    "raw_source_port": source_port,
                    "raw_target_port": target_port_raw,
                }
                out.append(normalized)

    write_jsonl(output_path, out)
    return out

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--connections", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()
    rows = normalize_connections(args.connections, args.output)
    print(f"[OK] normalized connections: {len(rows)} -> {args.output}")

if __name__ == "__main__":
    main()
