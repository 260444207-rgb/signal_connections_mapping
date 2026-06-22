#!/usr/bin/env python3
# -*- coding: utf-8 -*-

from __future__ import annotations

import argparse
from pathlib import Path
from typing import Dict, Any, List

from common import iter_jsonl, ensure_dir, FINAL_HEADERS
from generate_net_name import generate_net_name

SKIP_SHEETS = {"BLOCK_INFO", "LINK_INFO", "链路信息", "说明", "README", "INDEX", "目录"}
MULTI_RESULT_CONNECTION_ID_SEPARATOR = "#"

def decision_selected_pins(decision: Dict[str, Any]) -> List[str]:
    pins = decision.get("selected_pins")
    if isinstance(pins, list) and pins:
        return [str(pin).strip() for pin in pins if str(pin).strip()]
    pin = str(decision.get("selected_pin", "") or "").strip()
    return [pin] if pin else []

def decision_list_value(decision: Dict[str, Any], key: str, index: int, default: str = "") -> str:
    values = decision.get(key)
    if isinstance(values, list) and index < len(values):
        return str(values[index] or "")
    return str(decision.get(key[:-1] if key.endswith("s") else key, default) or default)

def is_placeholder_net_name(value: str) -> bool:
    """LINE/line 类名称是画图工具默认连线名，不能作为有效网络名。"""
    text = str(value or "").strip()
    return bool(text) and "line" in text.lower()

def is_invalid_model_net_name(value: str) -> bool:
    text = str(value or "").strip().upper()
    if not text:
        return True
    if is_placeholder_net_name(text):
        return True
    return text in {
        "INPUT",
        "OUTPUT",
        "IN",
        "OUT",
        "HIGH",
        "MEDIUM",
        "LOW",
        "MODEL_RESOLVED",
        "UNRESOLVED",
    }

def final_net_name(decision: Dict[str, Any], normalized: Dict[str, Any], index: int, selected_pin: str) -> str:
    """
    生成最终网络名。
    硬约束：没有原理图 pin 时网络名必须为空；LINE 类默认连线名不能覆盖网络命名。
    """
    if not str(selected_pin or "").strip():
        return ""

    connection_name = str(normalized.get("connection_name", "") or "").strip()
    if connection_name and not is_placeholder_net_name(connection_name):
        return generate_net_name(normalized, selected_pin)

    decision_net_name = decision_list_value(decision, "net_names", index)
    if decision_net_name and not is_invalid_model_net_name(decision_net_name):
        return decision_net_name

    return generate_net_name(normalized, selected_pin)

def expanded_connection_id(base_connection_id: str, original_connection_id: str, should_expand: bool, index: int) -> str:
    if not should_expand:
        return original_connection_id
    base = str(base_connection_id or original_connection_id or "").strip()
    return f"{base}{MULTI_RESULT_CONNECTION_ID_SEPARATOR}{index + 1}" if base else str(index + 1)

def normalize_selected_pins_for_row(normalized: Dict[str, Any], selected_pins: List[str]) -> tuple[List[str], int]:
    shape_info = normalized.get("signal_shape_info", {})
    if not isinstance(shape_info, dict) or not shape_info.get("is_expanded_member"):
        return selected_pins, 0
    if len(selected_pins) <= 1:
        return selected_pins, 0
    member_index = int(shape_info.get("member_index", 1) or 1) - 1
    member_index = max(0, min(member_index, len(selected_pins) - 1))
    return [selected_pins[member_index]], member_index

def inferred_parent_line_id(line_id: str) -> str:
    text = str(line_id or "")
    if "#" not in text:
        return ""
    parent, suffix = text.rsplit("#", 1)
    return parent if parent and suffix.isdigit() else ""

def decision_parent_line_id(decision: Dict[str, Any]) -> str:
    return str(decision.get("parent_line_id", "") or "").strip() or inferred_parent_line_id(decision.get("line_id", ""))

def child_decision_index(decision: Dict[str, Any]) -> int:
    line_id = str(decision.get("line_id", "") or "")
    if "#" not in line_id:
        return 0
    suffix = line_id.rsplit("#", 1)[1]
    return int(suffix) - 1 if suffix.isdigit() else 0

def display_connection_id_for_decision(normalized: Dict[str, Any], decision: Dict[str, Any]) -> str:
    parent = decision_parent_line_id(decision)
    line_id = str(decision.get("line_id", "") or "")
    if parent and line_id.startswith(parent + "#"):
        suffix = line_id.rsplit("#", 1)[1]
        base = str(normalized.get("base_connection_id") or normalized.get("connection_id") or "").strip()
        return f"{base}#{suffix}" if base else line_id
    return normalized.get("connection_id", "")

def group_child_decisions(decisions: Dict[str, Dict[str, Any]], normalized_ids: set[str]) -> Dict[str, List[Dict[str, Any]]]:
    grouped: Dict[str, List[Dict[str, Any]]] = {}
    for decision in decisions.values():
        line_id = decision.get("line_id", "")
        if line_id in normalized_ids:
            continue
        parent = decision_parent_line_id(decision)
        if parent:
            grouped.setdefault(parent, []).append(decision)
    for items in grouped.values():
        items.sort(key=lambda item: child_decision_index(item))
    return grouped

def rows_for_decision(normalized: Dict[str, Any], decision: Dict[str, Any], connection_id: str | None = None) -> List[Dict[str, Any]]:
    connection_name = normalized.get("connection_name", "") or ""
    selected_pins, decision_index_offset = normalize_selected_pins_for_row(normalized, decision_selected_pins(decision) or [""])
    should_expand = len(selected_pins) > 1
    base_connection_id = normalized.get("base_connection_id") or normalized.get("connection_id", "")
    rows: List[Dict[str, Any]] = []
    for idx, selected_pin in enumerate(selected_pins):
        decision_index = decision_index_offset + idx
        net_name = final_net_name(decision, normalized, decision_index, selected_pin)
        rendered_connection_id = connection_id or expanded_connection_id(
            base_connection_id,
            normalized.get("connection_id", ""),
            should_expand,
            idx,
        )
        rows.append({
            "源Block标识": normalized.get("source_block_id", ""),
            "源Block名称": normalized.get("source_block_name", ""),
            "源Port": normalized.get("source_port", ""),
            "目的Block标识": normalized.get("target_block_id", ""),
            "目的Block名称": normalized.get("target_block_name", ""),
            "目的Port": normalized.get("target_port", ""),
            "连线ID": rendered_connection_id,
            "连线名称": connection_name,
            "连线方向": normalized.get("direction", ""),
            "原理图Pin脚": selected_pin,
            "分析说明": decision_list_value(decision, "analyses", decision_index, decision.get("analysis", "")),
            "映射置信度": decision_list_value(decision, "confidences", decision_index, decision.get("confidence", "")),
            "网络命名": net_name,
        })
    return rows

def build_final_rows_by_template_sheet(normalized_path: str | Path, decisions_path: str | Path) -> Dict[str, List[Dict[str, Any]]]:
    """
    按 normalized_connection.output_sheet_name 分组。
    注意：output_sheet_name 来自输入框图 Excel 的原 sheet 名，不允许按 block_id/group 自行分组。
    """
    decisions = {d["line_id"]: d for d in iter_jsonl(decisions_path)}
    normalized_rows = list(iter_jsonl(normalized_path))
    normalized_ids = {row["line_id"] for row in normalized_rows}
    child_decisions = group_child_decisions(decisions, normalized_ids)
    grouped: Dict[str, List[Dict[str, Any]]] = {}

    for n in normalized_rows:
        sheet_name = n.get("output_sheet_name") or n.get("source_sheet_name")
        if not sheet_name:
            raise ValueError(f"normalized row missing output_sheet_name/source_sheet_name: {n.get('line_id')}")
        children = child_decisions.get(n["line_id"], [])
        if children:
            for child in children:
                grouped.setdefault(sheet_name, []).extend(rows_for_decision(n, child, display_connection_id_for_decision(n, child)))
            continue
        grouped.setdefault(sheet_name, []).extend(rows_for_decision(n, decisions.get(n["line_id"], {})))

    return grouped

def clear_sheet_values(ws):
    # 保留 sheet 本身，删除旧内容后从第一行重建标准结果表。
    if ws.max_row:
        ws.delete_rows(1, ws.max_row)

def write_rows(ws, rows: List[Dict[str, Any]]):
    clear_sheet_values(ws)
    ws.append(FINAL_HEADERS)
    for row in rows:
        ws.append([row.get(h, "") for h in FINAL_HEADERS])

    # 简单格式，避免强依赖复杂样式
    try:
        from openpyxl.styles import Font, Alignment
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            ws.column_dimensions[col[0].column_letter].width = min(40, max(10, max_len + 2))
    except Exception:
        pass

def render_template_sheets(template_excel: str | Path, normalized_path: str | Path, decisions_path: str | Path, output_excel: str | Path) -> Dict[str, Any]:
    """
    正式信号接口列表输出：
    1. 复制输入框图 Excel 的 sheet 结构。
    2. 只向输入 Excel 已存在的器件 sheet 写入分析结果。
    3. 不根据 block_id/group/source_sheet 自行创建新 sheet。
    4. 不把所有结果合并到一个总 sheet。
    """
    from openpyxl import load_workbook

    template_excel = Path(template_excel)
    output_excel = Path(output_excel)
    ensure_dir(output_excel.parent)

    wb = load_workbook(template_excel)
    grouped = build_final_rows_by_template_sheet(normalized_path, decisions_path)

    input_sheets = set(wb.sheetnames)
    missing_in_template = sorted([s for s in grouped.keys() if s not in input_sheets])

    # 严格模式：不能创建新 sheet，避免“自己分组输出”
    if missing_in_template:
        raise ValueError(
            "以下输出 sheet 不存在于输入框图表格中，禁止自动创建新 sheet："
            + ", ".join(missing_in_template)
        )

    written_sheets = []
    for sheet_name in wb.sheetnames:
        if sheet_name.upper() in SKIP_SHEETS:
            written_sheets.append({"sheet_name": sheet_name, "row_count": "preserved"})
            continue
        rows = grouped.get(sheet_name, [])
        # 输入有这个器件 sheet，就保留该 sheet；有分析结果则写入，无结果则写入空表头。
        ws = wb[sheet_name]
        write_rows(ws, rows)
        written_sheets.append({"sheet_name": sheet_name, "row_count": len(rows)})

    wb.save(output_excel)
    wb.close()

    return {
        "output_excel": str(output_excel),
        "written_sheets": written_sheets,
        "missing_in_template": missing_in_template,
    }

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--template-excel", required=True, help="输入框图 Excel，输出必须严格沿用它的 sheet 分页")
    parser.add_argument("--normalized", required=True)
    parser.add_argument("--decisions", required=True)
    parser.add_argument("--output-excel", required=True)
    args = parser.parse_args()

    result = render_template_sheets(args.template_excel, args.normalized, args.decisions, args.output_excel)
    print(f"[OK] rendered template sheets -> {result['output_excel']}")
    for item in result["written_sheets"]:
        print(f"  - {item['sheet_name']}: {item['row_count']} rows")

if __name__ == "__main__":
    main()
