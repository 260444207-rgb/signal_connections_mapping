from __future__ import annotations

import sys
import tempfile
import unittest
import os
import json
from pathlib import Path
from unittest import mock


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from signal_mapping.build_model_resolution_tasks import extract_rule_blocks, match_rule_sections
from signal_mapping.check_subagent_outputs import check_subagent_outputs
from signal_mapping.common import (
    FINAL_HEADERS,
    extract_component_uuid,
    iter_jsonl,
    load_pin_catalog,
    normalize_component_reference,
    pins_for_part,
    worksheet_header_map,
    write_jsonl,
)
from signal_mapping.render_template_sheets import build_final_rows_by_template_sheet, render_template_sheets, rows_for_decision
from signal_mapping.normalize_connections import normalize_connections
from signal_mapping.validate_mapping import validate_mapping
from signal_mapping.infer_signal_shapes import (
    build_companion_wire_index,
    expanded_rows_for_signal_shape,
    infer_signal_shapes,
    infer_signal_shape_for_row,
    iter_rule_sections,
    matching_rule_hint,
)
from signal_mapping.route_model_resolution import route_model_resolution
from signal_mapping.build_model_resolution_tasks import (
    PIN_GROUP_THRESHOLD,
    build_diagram_link_context,
    build_link_family_profiles,
    build_pin_allocation_context,
    build_pin_group_catalog,
    compact_normalized_connection,
    pins_for_groups,
    render_shared_prompt,
    session_link_family_profile,
)
from run_pipeline import (
    PIPELINE_CONFIG_FILENAME,
    FINISH_COMMAND_FILENAME,
    append_rule_path,
    build_combined_rules,
    combined_rules_are_current,
    discover_external_device_rules,
    load_pipeline_run_config,
    run_finish,
    save_pipeline_run_config,
)
from run_pipeline import resolve_signal_interface_task_dir


def make_group(part: str, link_family: str, mapping_family: str, source: str, target: str) -> dict:
    return {
        "source_device_signature": f"DEVICE_INFO:{part}",
        "link_family_ids": [link_family],
        "mapping_families": [mapping_family],
        "target_device_signatures": [f"TARGET_CONTEXT:{target}"],
        "source_device_instances": [source],
        "target_device_instances": [target],
        "user_link_infos": [],
        "device_role_infos": [],
    }


def make_row(
    part: str,
    source_port: str,
    target: str,
    target_port: str,
    link_family: str = "",
    source: str = "SRC",
    shape: str = "scalar",
) -> dict:
    return {
        "source_part_id": part,
        "source_block_name": source,
        "source_block_id": source,
        "source_port": source_port,
        "target_block_name": target,
        "target_port": target_port,
        "connection_name": "LINE_1",
        "link_family_id": link_family,
        "link_instance_id": "",
        "user_link_info": "",
        "device_role_info": "",
        "signal_shape": shape,
        "signal_shape_info": {"shape": shape},
    }


class LayeredRuleRecallTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.task_dir = Path(self.temp_dir.name)
        self.combined_path = build_combined_rules(self.task_dir)
        self.rule_text = self.combined_path.read_text(encoding="utf-8")
        self.blocks = extract_rule_blocks(self.rule_text)

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def test_rule_bus_with_explicit_width_expands_before_mapping(self) -> None:
        row = {
            "line_id": "L_BUS",
            "source_part_id": "DEV_A",
            "source_block_name": "DEV_A",
            "source_port": "SPI_DATA",
            "target_port": "DATA",
            "connection_name": "SPI_DATA",
            "connection_id": "C_BUS",
            "base_connection_id": "C_BUS",
            "expansion_count": 1,
        }
        rule = """### RULE: DEV_A SPI_DATA 总线规则
适用条件：DEV_A / SPI_DATA
SPI_DATA 声明为 4-bit 总线，实际成员为 D0、D1、D2。
"""

        info = infer_signal_shape_for_row(row, ["D0", "D1", "D2"], rule)
        expanded = expanded_rows_for_signal_shape(row, info)

        self.assertEqual("bus", info["shape"])
        self.assertEqual(3, info["expected_physical_pin_count"])
        self.assertEqual(4, info["evidence"]["matched_rule_shape_hint"]["declared_rule_bus_width"])
        self.assertEqual(["D0", "D1", "D2"], info["evidence"]["matched_rule_shape_hint"]["matched_bus_pins"])
        self.assertFalse(info["requires_model_expansion"])
        self.assertEqual(
            ["L_BUS#1", "L_BUS#2", "L_BUS#3"],
            [member["line_id"] for member in expanded],
        )
        self.assertTrue(all(member["signal_shape_info"]["is_expanded_member"] for member in expanded))

    def test_rule_bus_without_width_requires_model_children(self) -> None:
        row = {
            "line_id": "L_SPI",
            "source_part_id": "DEV_A",
            "source_block_name": "DEV_A",
            "source_port": "SPI_DATA",
            "target_port": "SPI",
            "connection_name": "SPI_DATA",
            "expansion_count": 1,
        }
        rule = """### RULE: DEV_A SPI_DATA 总线规则
适用条件：DEV_A / SPI_DATA
SPI_DATA 是总线，但没有列出实际 pin 名称。
"""

        info = infer_signal_shape_for_row(row, ["PIN_A", "PIN_B", "PIN_C", "PIN_D"], rule)
        compact = compact_normalized_connection({**row, "signal_shape": "bus", "signal_shape_info": info})

        self.assertEqual("bus", info["shape"])
        self.assertEqual("SPI_DATA", info["bus_name"])
        self.assertTrue(info["requires_model_expansion"])
        self.assertFalse(info["expected_physical_pin_count_known"])
        self.assertTrue(info["needs_model_shape_review"])
        self.assertEqual("model_must_emit_parent_line_id_children", info["line_id_expansion_policy"])
        self.assertEqual("SPI_DATA", compact["signal_shape_info"]["bus_name"])
        self.assertTrue(compact["signal_shape_info"]["requires_model_expansion"])

    def test_bus_rule_body_is_ignored_when_rule_title_names_another_device(self) -> None:
        row = {
            "line_id": "L_WRONG_DEVICE",
            "source_part_id": "DEV_A",
            "source_block_name": "DEV_A",
            "source_port": "SPI_DATA",
            "target_port": "SPI",
            "connection_name": "SPI_DATA",
            "expansion_count": 1,
        }
        rule = """### RULE: DEV_B SPI_DATA 总线规则
适用条件：DEV_A / SPI_DATA
SPI_DATA 是 4-bit 总线，成员为 D0、D1、D2、D3。
"""

        info = infer_signal_shape_for_row(row, ["D0", "D1", "D2", "D3"], rule)

        self.assertEqual("scalar", info["shape"])
        self.assertNotIn("matched_rule_shape_hint", info["evidence"])

    def test_preexpanded_bus_member_keeps_bus_metadata_without_double_expansion(self) -> None:
        row = {
            "line_id": "L_PRE#2",
            "base_line_id": "L_PRE",
            "source_port": "DATA[3:0]",
            "target_port": "DATA",
            "connection_id": "C_PRE#2",
            "base_connection_id": "C_PRE",
            "expansion_index": 2,
            "expansion_count": 4,
        }

        info = infer_signal_shape_for_row(row, ["D0", "D1", "D2", "D3"], "")
        expanded = expanded_rows_for_signal_shape(row, info)
        compact = compact_normalized_connection({**row, "signal_shape": "bus", "signal_shape_info": info})

        self.assertEqual("bus", info["shape"])
        self.assertTrue(info["is_expanded_member"])
        self.assertEqual("L_PRE", info["parent_line_id"])
        self.assertEqual(2, info["member_index"])
        self.assertEqual(4, info["member_count"])
        self.assertEqual(1, info["expected_physical_pin_count"])
        self.assertEqual(["L_PRE#2"], [member["line_id"] for member in expanded])
        self.assertTrue(compact["signal_shape_info"]["is_expanded_member"])
        self.assertEqual(4, compact["signal_shape_info"]["parent_expected_physical_pin_count"])

    def test_normalize_keeps_explicit_bus_as_one_row_until_rule_and_pin_validation(self) -> None:
        connections = self.task_dir / "connections.csv"
        normalized = self.task_dir / "normalized.jsonl"
        connections.write_text(
            "source_port,target_port,connection_id\nDATA[3:0],DATA,C_DATA\n",
            encoding="utf-8",
        )

        rows = normalize_connections(connections, normalized)

        self.assertEqual(1, len(rows))
        self.assertEqual(0, rows[0]["source_row_number"])
        self.assertEqual(1, rows[0]["expansion_count"])
        self.assertEqual(4, rows[0]["declared_bus_width"])
        self.assertEqual("C_DATA", rows[0]["connection_id"])

    def test_scheme_e_companion_pair_rows_are_not_expanded_again(self) -> None:
        normalized = self.task_dir / "normalized_diff.jsonl"
        pins = self.task_dir / "pins_diff.json"
        output = self.task_dir / "shaped_diff.jsonl"
        report = self.task_dir / "shape_report.jsonl"
        rows = [
            {
                "line_id": "L_N",
                "source_sheet_name": "RF_SHEET",
                "source_part_id": "PART_RF",
                "source_block_id": "CHA",
                "source_port": "CHA_RF_IN_N",
                "target_block_id": "AFE",
                "target_port": "AFE1_IOUTP_TXM3",
                "connection_id": "100677",
                "connection_name": "100677",
            },
            {
                "line_id": "L_P",
                "source_sheet_name": "RF_SHEET",
                "source_part_id": "PART_RF",
                "source_block_id": "CHA",
                "source_port": "CHA_RF_IN_P",
                "target_block_id": "AFE",
                "target_port": "AFE1_IOUTN_TXM3",
                "connection_id": "100679",
                "connection_name": "100679",
            },
        ]
        write_jsonl(normalized, rows)
        pins.write_text(
            json.dumps({"PART_RF": ["CHA_RF_IN_N", "CHA_RF_IN_P"]}),
            encoding="utf-8",
        )

        shaped = infer_signal_shapes(normalized, pins, output, report)

        self.assertEqual(2, len(shaped))
        self.assertEqual(["L_N", "L_P"], [row["line_id"] for row in shaped])
        self.assertTrue(all(row["signal_shape"] == "scalar" for row in shaped))
        self.assertTrue(all(
            row["signal_shape_info"]["evidence"]["differential_member_guard"]["companion_wire"]
            for row in shaped
        ))

    def test_scheme_e_companion_detection_allows_reversed_block_direction(self) -> None:
        rows = [
            {
                "line_id": "L_FORWARD",
                "source_sheet_name": "SHEET_A",
                "source_block_id": "BLOCK_A",
                "source_port": "LINK_N",
                "target_block_id": "BLOCK_B",
                "target_port": "RX_N",
            },
            {
                "line_id": "L_REVERSED",
                "source_sheet_name": "SHEET_A",
                "source_block_id": "BLOCK_B",
                "source_port": "RX_P",
                "target_block_id": "BLOCK_A",
                "target_port": "LINK_P",
            },
        ]

        companions = build_companion_wire_index(rows)

        self.assertEqual("L_REVERSED", companions["L_FORWARD"]["companion_line_id"])
        self.assertEqual("reversed", companions["L_FORWARD"]["direction_relation"])
        self.assertEqual("L_FORWARD", companions["L_REVERSED"]["companion_line_id"])

    def test_scheme_e_active_low_and_power_target_are_scalar(self) -> None:
        active_low = {
            "line_id": "L_POR",
            "source_port": "POR_N",
            "target_port": "RESET2",
            "connection_name": "POR reset",
        }
        power = {
            "line_id": "L_PWR",
            "source_port": "LNASIP_RXD_OUT_N",
            "target_port": "1V8",
            "connection_name": "power connection",
        }

        active_info = infer_signal_shape_for_row(active_low, [], "")
        power_info = infer_signal_shape_for_row(power, [], "")

        self.assertEqual("scalar", active_info["shape"])
        self.assertFalse(active_info["needs_model_shape_review"])
        self.assertTrue(active_info["evidence"]["differential_member_guard"]["low_active_signal"])
        self.assertEqual("scalar", power_info["shape"])
        self.assertFalse(power_info["needs_model_shape_review"])
        self.assertTrue(power_info["evidence"]["differential_member_guard"]["power_target"])

    def test_scheme_e_unpaired_polarity_port_requires_review_without_expansion(self) -> None:
        row = {
            "line_id": "L_MEMBER",
            "source_port": "RFIN_N",
            "target_port": "AFE_INPUT",
            "connection_name": "RF input",
        }

        info = infer_signal_shape_for_row(row, ["RFIN_P", "RFIN_N"], "")
        expanded = expanded_rows_for_signal_shape(row, info)

        self.assertEqual("scalar", info["shape"])
        self.assertTrue(info["needs_model_shape_review"])
        self.assertEqual(["L_MEMBER"], [item["line_id"] for item in expanded])

    def test_scheme_e_connection_level_marker_still_expands_whole_pair(self) -> None:
        row = {
            "line_id": "L_DIFF",
            "source_port": "DAC00",
            "target_port": "RFIN0",
            "connection_name": "TX differential pair",
            "connection_id": "C_DIFF",
        }

        info = infer_signal_shape_for_row(row, ["TX_AFE0_00_P", "TX_AFE0_00_N"], "")
        expanded = expanded_rows_for_signal_shape(row, info)

        self.assertEqual("differential", info["shape"])
        self.assertEqual(["L_DIFF#1", "L_DIFF#2"], [item["line_id"] for item in expanded])

    def test_scheme_e_function_and_rule_hints_do_not_auto_expand(self) -> None:
        function_row = {
            "line_id": "L_CLKREF",
            "source_port": "AFE_CLKREF",
            "target_port": "CLKOUT1",
            "connection_name": "clock reference",
        }
        rule_row = {
            "line_id": "L_CPRI",
            "source_part_id": "DEV_CPRI",
            "source_block_name": "DEV_CPRI",
            "source_port": "CPRI0_RX",
            "target_port": "default1",
            "connection_name": "CPRI receive",
        }
        rule = """### RULE: DEV_CPRI CPRI 接收规则
适用条件：DEV_CPRI / CPRI0_RX
CPRI0_RX 是差分接口，映射到 CPRI0_RX_P 和 CPRI0_RX_N。
"""

        function_info = infer_signal_shape_for_row(
            function_row,
            ["AFE_CLKREF_P", "AFE_CLKREF_N"],
            "",
        )
        rule_info = infer_signal_shape_for_row(
            rule_row,
            ["CPRI0_RX_P", "CPRI0_RX_N"],
            rule,
        )

        self.assertEqual("scalar", function_info["shape"])
        self.assertTrue(function_info["needs_model_shape_review"])
        self.assertEqual("scalar", rule_info["shape"])
        self.assertTrue(rule_info["needs_model_shape_review"])

    def test_session_line_examples_require_explicit_family_and_same_device(self) -> None:
        groups = [
            {
                "context_group_id": "CTX_A",
                "source_device_signature": "DEVICE_INFO:PART_A",
                "link_family_ids": ["SPI_CTRL"],
                "link_family_sources": ["explicit_link_info", "fallback_mapping_family"],
                "line_ids": ["A_EXPLICIT", "A_INFERRED", "A_AMBIGUOUS"],
            },
            {
                "context_group_id": "CTX_B",
                "source_device_signature": "DEVICE_INFO:PART_B",
                "link_family_ids": ["SPI_CTRL"],
                "link_family_sources": ["explicit_link_info"],
                "line_ids": ["B_EXPLICIT"],
            },
        ]
        rows = {
            "A_EXPLICIT": {
                "source_port": "SPI_CLK",
                "target_port": "CLK",
                "link_family_id": "SPI_CTRL",
            },
            "A_INFERRED": {
                "source_port": "SPI_CLK",
                "target_port": "CLK",
            },
            "A_AMBIGUOUS": {
                "source_port": "SPI_CLK",
                "target_port": "CLK",
                "link_contexts": [{"link_family_id": "SPI_CTRL"}],
            },
            "B_EXPLICIT": {
                "source_port": "SPI_CLK",
                "target_port": "CLK",
                "link_family_id": "SPI_CTRL",
            },
        }

        profile = build_link_family_profiles(groups, rows)["SPI_CTRL"]
        self.assertEqual(
            ["A_EXPLICIT", "B_EXPLICIT"],
            [example["line_id"] for example in profile["line_examples"]],
        )
        self.assertEqual(
            ["A_AMBIGUOUS"],
            [example["line_id"] for example in profile["ambiguous_sheet_member_line_examples"]],
        )

        scoped = session_link_family_profile(profile, "DEVICE_INFO:PART_A")
        self.assertEqual(
            ["A_EXPLICIT"],
            [example["line_id"] for example in scoped["line_examples"]],
        )
        self.assertEqual([], scoped["ambiguous_sheet_member_line_examples"])

    def test_collects_all_layered_rule_files_with_metadata(self) -> None:
        self.assertEqual(18, len(self.blocks))
        by_id = {block["rule_id"]: block for block in self.blocks}
        self.assertEqual("device", by_id["0302078562"]["layer"])
        self.assertEqual("link", by_id["LINK_RF_TX"]["layer"])
        self.assertEqual("signal", by_id["SIGNAL_DIFF"]["layer"])
        self.assertEqual("global", by_id["GLOBAL_MAPPING"]["layer"])
        self.assertNotIn("rf_tx_chain.md", by_id["LINK_POWER"]["text"].lower())

    def test_shape_rule_sections_are_not_truncated(self) -> None:
        sections = iter_rule_sections(self.rule_text)
        self.assertEqual(18, len(sections))
        self.assertTrue(all(section.startswith("### RULE:") for section in sections))
        self.assertTrue(all(len(section) > 20 for section in sections))

    def test_shape_hint_is_local_to_matching_subsection(self) -> None:
        control_row = make_row(
            "302078562", "TX_SW0", "TXVGA00", "EN_CHA", "RF_TX_CHAIN", "SROC"
        )
        control_hint = matching_rule_hint(control_row, self.rule_text)
        self.assertNotIn("differential", control_hint.get("shape_hints", []))
        self.assertFalse(control_hint.get("strong_shape_hint", False))

        differential_row = make_row(
            "47151290", "RFIN0", "SROC", "DAC00", "RF_TX_CHAIN", "TXVGA"
        )
        differential_hint = matching_rule_hint(differential_row, self.rule_text)
        self.assertIn("differential", differential_hint["shape_hints"])
        self.assertTrue(differential_hint["strong_shape_hint"])

    def test_recall_obeys_link_device_signal_global_layers(self) -> None:
        group = make_group("47151290", "RF_TX_CHAIN", "RF_CHAIN", "TXVGA", "SROC")
        row = make_row(
            "47151290", "RFIN0", "SROC", "DAC00", "RF_TX_CHAIN", "TXVGA", "differential"
        )
        matches = match_rule_sections(self.blocks, group, [row])
        ids = [match["rule_id"] for match in matches]
        self.assertEqual(
            ["LINK_RF_TX", "47151290", "SIGNAL_DIFF", "SIGNAL_SPECIAL", "GLOBAL_MAPPING"],
            ids,
        )
        self.assertEqual("exact_link_family", matches[0]["match_type"])
        self.assertEqual("exact_source_part", matches[1]["match_type"])

    def test_sroc_part_id_exact_match_ignores_leading_zero(self) -> None:
        group = make_group("302078562", "RF_TX_CHAIN", "RF_CHAIN", "SROC", "TXVGA00")
        row = make_row(
            "302078562", "TX_SW0", "TXVGA00", "EN_CHA", "RF_TX_CHAIN", "SROC"
        )
        matches = match_rule_sections(self.blocks, group, [row])
        sroc = next(match for match in matches if match["rule_id"] == "0302078562")
        self.assertEqual("exact_source_part", sroc["match_type"])
        self.assertNotIn("AMC7964", [match["rule_id"] for match in matches])
        self.assertNotIn("TXCAL_RXCAL", [match["rule_id"] for match in matches])

    def test_device_rule_recall_rejects_body_only_device_match(self) -> None:
        text = """<!-- SOURCE: rules/device_rules/wrong.md -->
### RULE: DEV_B 错误器件规则
适用条件：
源端器件：DEV_A
DEV_A 的 SPI_DATA 是总线。
"""
        blocks = extract_rule_blocks(text)
        group = make_group("DEV_A", "SPI_CTRL", "SPI_CTRL", "DEV_A", "TARGET")
        row = make_row("DEV_A", "SPI_DATA", "TARGET", "SPI", "SPI_CTRL", "DEV_A")

        matches = match_rule_sections(blocks, group, [row])

        self.assertEqual([], matches)

    def test_extracts_component_uuid_from_block_info_metadata(self) -> None:
        component_uuid = "6c33dd64-9afd-4e31-87c4-5023b323be7f"
        samples = [
            json.dumps({"componentName": "SROC", "componentUuid": component_uuid}),
            rf'prefix {{\"componentUuid\":\"{component_uuid}\",\"name\":\"SROC\"}} suffix',
            f"componentUuid={component_uuid};componentName=SROC",
            json.dumps(json.dumps({"componentUuid": component_uuid})),
        ]
        for sample in samples:
            with self.subTest(sample=sample):
                self.assertEqual(component_uuid, extract_component_uuid(sample))
                self.assertEqual(component_uuid, normalize_component_reference(sample))

    def test_loads_component_uuid_pin_catalog_record_format(self) -> None:
        component_uuid = "6c33dd64-9afd-4e31-87c4-5023b323be7f"
        pins_path = self.task_dir / "component_pins.json"
        pins_path.write_text(
            json.dumps({
                "components": [
                    {
                        "componentUuid": component_uuid,
                        "pins": [
                            {"pinName": "PA_SW0"},
                            {"name": "FEM_TDDSW00"},
                        ],
                    }
                ]
            }),
            encoding="utf-8",
        )

        catalog = load_pin_catalog(pins_path)

        self.assertEqual(["PA_SW0", "FEM_TDDSW00"], catalog[component_uuid])
        long_block_info_value = json.dumps({"componentUuid": component_uuid, "other": "metadata"})
        self.assertEqual(["PA_SW0", "FEM_TDDSW00"], pins_for_part(catalog, long_block_info_value))

    def test_loads_component_uuid_one_pin_per_record_format(self) -> None:
        component_uuid = "6c33dd64-9afd-4e31-87c4-5023b323be7f"
        pins_path = self.task_dir / "component_pin_rows.json"
        pins_path.write_text(
            json.dumps([
                {"componentUuid": component_uuid, "pinName": "PIN_A"},
                {"componentUuid": component_uuid, "pinInfo": {"pinName": "PIN_B"}},
            ]),
            encoding="utf-8",
        )

        catalog = load_pin_catalog(pins_path)

        self.assertEqual(["PIN_A", "PIN_B"], catalog[component_uuid])

    def test_preserves_legacy_part_number_pin_catalog(self) -> None:
        pins_path = self.task_dir / "legacy_pins.json"
        pins_path.write_text(
            json.dumps({"0302078562": ["PIN_A", "PIN_B"]}),
            encoding="utf-8",
        )

        catalog = load_pin_catalog(pins_path)

        self.assertEqual(["PIN_A", "PIN_B"], pins_for_part(catalog, "302078562"))

    def test_general_rules_are_deterministically_recalled(self) -> None:
        power_group = make_group("999999", "POWER_CHAIN", "POWER_ENABLE", "MYSTERY", "PMU")
        power_row = make_row("999999", "VDD_3V3", "PMU", "VOUT", "POWER_CHAIN", "MYSTERY")
        power_ids = [
            match["rule_id"]
            for match in match_rule_sections(self.blocks, power_group, [power_row])
        ]
        self.assertIn("LINK_POWER", power_ids)
        self.assertIn("SIGNAL_POWER", power_ids)
        self.assertIn("SIGNAL_SPECIAL", power_ids)
        self.assertIn("GLOBAL_MAPPING", power_ids)
        self.assertNotIn("47151290", power_ids)

        diff_group = make_group(
            "999999", "LOCAL_DEVICE_MAPPING", "DIFFERENTIAL_PAIR", "MYSTERY", "ADC"
        )
        diff_row = make_row("999999", "ADC0", "ADC", "RFIN0", source="MYSTERY", shape="differential")
        diff_ids = [
            match["rule_id"]
            for match in match_rule_sections(self.blocks, diff_group, [diff_row])
        ]
        self.assertEqual(["SIGNAL_DIFF", "SIGNAL_SPECIAL", "GLOBAL_MAPPING"], diff_ids)

    def test_rule_snapshot_detects_external_rule_changes(self) -> None:
        user_rule = self.task_dir / "user_rules.md"
        user_rule.write_text("### RULE: USER_A 用户规则\n适用条件：\n通用\n", encoding="utf-8")
        build_combined_rules(self.task_dir, user_rules=str(user_rule))
        self.assertTrue(combined_rules_are_current(self.task_dir, user_rules=str(user_rule)))
        user_rule.write_text("### RULE: USER_B 更新规则\n适用条件：\n通用\n", encoding="utf-8")
        self.assertFalse(combined_rules_are_current(self.task_dir, user_rules=str(user_rule)))

    def test_auto_discovers_external_device_rules_file(self) -> None:
        external_rule = self.task_dir / "design" / "external_device_rules.md"
        external_rule.parent.mkdir()
        external_rule.write_text(
            "### RULE: EXT_DEVICE_API 接口器件规则\n"
            "适用条件：\n"
            "源端器件：EXT_DEV / 12345\n"
            "规则摘要：接口返回规则。\n",
            encoding="utf-8",
        )
        discovered = discover_external_device_rules(self.task_dir, "")
        self.assertEqual(external_rule, discovered)

        project_rules = append_rule_path("", discovered)
        combined = build_combined_rules(self.task_dir, project_rules=project_rules)
        text = combined.read_text(encoding="utf-8")
        self.assertIn("EXT_DEVICE_API", text)
        self.assertIn("接口返回规则", text)

    def test_project_rules_accept_multiple_paths(self) -> None:
        rule_a = self.task_dir / "rule_a.md"
        rule_b = self.task_dir / "rule_b.md"
        rule_a.write_text("### RULE: EXT_A 外部规则A\n适用条件：\n通用\n", encoding="utf-8")
        rule_b.write_text("### RULE: EXT_B 外部规则B\n适用条件：\n通用\n", encoding="utf-8")

        project_rules = append_rule_path(str(rule_a), rule_b)
        self.assertIn(os.pathsep, project_rules)

        combined = build_combined_rules(self.task_dir, project_rules=project_rules)
        text = combined.read_text(encoding="utf-8")
        self.assertIn("EXT_A", text)
        self.assertIn("EXT_B", text)

        project_rules_again = append_rule_path(project_rules, rule_b)
        self.assertEqual(project_rules, project_rules_again)

    def test_resolves_signal_interface_task_subdir(self) -> None:
        root = self.task_dir / "data_uuid"
        self.assertEqual(root / "signal_interface", resolve_signal_interface_task_dir(root))
        already = root / "signal_interface"
        self.assertEqual(already, resolve_signal_interface_task_dir(already))

    def test_external_rule_overrides_same_builtin_rule_id(self) -> None:
        user_rule = self.task_dir / "rules" / "user_rules.md"
        user_rule.parent.mkdir()
        user_rule.write_text(
            "### RULE: 47151290 用户覆盖 TXVGA\n"
            "适用条件：\n"
            "源端器件：TXVGA / 47151290\n"
            "规则摘要：用户覆盖内容。\n",
            encoding="utf-8",
        )
        combined = build_combined_rules(self.task_dir, user_rules=str(user_rule))
        blocks = extract_rule_blocks(combined.read_text(encoding="utf-8"))
        group = make_group("47151290", "RF_TX_CHAIN", "RF_CHAIN", "TXVGA", "SROC")
        row = make_row("47151290", "RFIN0", "SROC", "DAC00", "RF_TX_CHAIN", "TXVGA")
        matches = match_rule_sections(blocks, group, [row])
        device_rule = next(match for match in matches if match["rule_id"] == "47151290")
        self.assertEqual("custom", device_rule["layer"])
        self.assertIn("用户覆盖内容", device_rule["text"])

    def test_spi_rule_uses_sequential_unallocated_pin_policy(self) -> None:
        spi_rule = (ROOT / "rules" / "signal_rules" / "spi_bus.md").read_text(encoding="utf-8")
        self.assertIn("尚未使用的合法 SPI pin 顺序分配", spi_rule)
        self.assertNotIn("不能唯一映射到具体 pin，需要人工确认", spi_rule)

    def test_model_routing_uses_pin_info_gate_without_pin_candidates(self) -> None:
        normalized = self.task_dir / "normalized.jsonl"
        pins = self.task_dir / "pins.json"
        decisions = self.task_dir / "decisions.jsonl"
        needs_model = self.task_dir / "needs_model.jsonl"
        rows = [
            {"line_id": "L1", "source_part_id": "1", "signal_shape": "scalar"},
            {"line_id": "L2", "source_part_id": "MISSING", "signal_shape": "scalar"},
        ]
        normalized.write_text(
            "\n".join(json.dumps(row, ensure_ascii=False) for row in rows) + "\n",
            encoding="utf-8",
        )
        pins.write_text(
            json.dumps({"001": ["PIN_A"]}, ensure_ascii=False),
            encoding="utf-8",
        )

        route_model_resolution(normalized, pins, decisions, needs_model)

        routed = [
            json.loads(line)
            for line in needs_model.read_text(encoding="utf-8").splitlines()
            if line.strip()
        ]
        routed_decisions = [
            json.loads(line)
            for line in decisions.read_text(encoding="utf-8").splitlines()
            if line.strip()
        ]
        self.assertEqual(["L1"], [row["line_id"] for row in routed])
        self.assertEqual(2, len(routed_decisions))
        self.assertTrue(all(row["selected_pin"] == "" for row in routed_decisions))

    def test_exact_source_port_pin_is_pre_resolved_without_subagent_task(self) -> None:
        normalized = self.task_dir / "normalized.jsonl"
        pins = self.task_dir / "pins.json"
        decisions = self.task_dir / "decisions.jsonl"
        needs_model = self.task_dir / "needs_model.jsonl"
        rows = [
            {
                "line_id": "L_DIRECT",
                "source_part_id": "001",
                "source_port": "PA_SW0",
                "signal_shape_info": {"shape": "scalar", "needs_model_shape_review": False},
            },
            {
                "line_id": "L_SEMANTIC",
                "source_part_id": "001",
                "source_port": "LOGICAL_CTRL",
                "signal_shape_info": {"shape": "scalar", "needs_model_shape_review": False},
            },
        ]
        normalized.write_text(
            "\n".join(json.dumps(row, ensure_ascii=False) for row in rows) + "\n",
            encoding="utf-8",
        )
        pins.write_text(
            json.dumps({"001": ["PA_SW0", "PA_PD_SW0"]}, ensure_ascii=False),
            encoding="utf-8",
        )

        route_model_resolution(normalized, pins, decisions, needs_model)

        routed = [json.loads(line) for line in needs_model.read_text(encoding="utf-8").splitlines() if line]
        by_line = {
            row["line_id"]: row
            for row in (json.loads(line) for line in decisions.read_text(encoding="utf-8").splitlines() if line)
        }
        self.assertEqual(["L_SEMANTIC"], [row["line_id"] for row in routed])
        self.assertEqual("PA_SW0", by_line["L_DIRECT"]["selected_pin"])
        self.assertEqual("pre_resolved", by_line["L_DIRECT"]["decision_type"])
        self.assertEqual("High", by_line["L_DIRECT"]["confidence"])
        self.assertTrue(all("net_name" not in row and "net_names" not in row for row in by_line.values()))

    def test_validation_rejects_second_mapping_of_exact_source_port_pin(self) -> None:
        normalized = self.task_dir / "normalized.jsonl"
        decisions = self.task_dir / "decisions.jsonl"
        pins = self.task_dir / "pins.json"
        report = self.task_dir / "validation.json"
        normalized.write_text(
            json.dumps({
                "line_id": "L1",
                "source_part_id": "001",
                "source_sheet_name": "SROC",
                "source_port": "PA_SW0",
                "signal_shape_info": {"shape": "scalar", "needs_model_shape_review": False},
            }) + "\n",
            encoding="utf-8",
        )
        decisions.write_text(
            json.dumps({
                "line_id": "L1",
                "selected_pin": "PA_PD_SW0",
                "decision_type": "model_resolved",
                "confidence": "High",
                "analysis": "incorrect second mapping",
            }) + "\n",
            encoding="utf-8",
        )
        pins.write_text(json.dumps({"001": ["PA_SW0", "PA_PD_SW0"]}), encoding="utf-8")

        result = validate_mapping(normalized, decisions, pins, report)

        self.assertEqual("ERROR", result["status"])
        self.assertTrue(any("must equal source_port" in error["message"] for error in result["errors"]))

    def test_subagent_payload_helpers_do_not_emit_pin_candidates(self) -> None:
        row = {
            "line_id": "L1",
            "source_sheet_name": "DEV0",
            "source_part_id": "001",
            "source_block_id": "B1",
            "source_block_name": "DEV",
            "source_port": "CTRL",
            "target_block_name": "LOAD",
            "target_port": "EN",
            "connection_id": "C1",
            "connection_name": "CTRL_NET",
            "signal_shape_info": {"shape": "scalar"},
        }
        payload = build_pin_allocation_context([row])
        serialized = json.dumps(payload, ensure_ascii=False)
        self.assertNotIn("candidate", serialized.lower())
        self.assertNotIn("rough_pin_search_hints", serialized)
        self.assertNotIn("line_pin_domains", serialized)
        self.assertNotIn("candidate_mappings", render_shared_prompt())
        self.assertIn("分析说明不得添加网络命名依据", render_shared_prompt())

    def test_shared_prompt_requires_model_subagent_with_30_minute_timeout(self) -> None:
        prompt = render_shared_prompt()
        self.assertIn("模型 subagent", prompt)
        self.assertIn("sessions_spawn", prompt)
        self.assertIn("30 分钟", prompt)
        self.assertIn("1800000 ms", prompt)
        self.assertIn("禁止编写 Python / PowerShell / JavaScript 等脚本", prompt)
        self.assertIn("selected_pin 的语义裁决必须由模型完成", prompt)
        self.assertIn("selected_pin 必须直接等于 source_port", prompt)

    def test_task_context_uses_one_canonical_per_line_record(self) -> None:
        row = {
            "line_id": "L1",
            "base_line_id": "L1",
            "source_sheet_name": "DEV0",
            "output_sheet_name": "DEV0",
            "source_part_id": "001",
            "source_block_id": "B1",
            "source_block_name": "DEV",
            "source_port": "CTRL",
            "raw_source_port": "CTRL",
            "target_block_id": "B2",
            "target_block_name": "LOAD",
            "target_port": "EN",
            "raw_target_port": "EN",
            "connection_id": "C1",
            "base_connection_id": "C1",
            "connection_name": "CTRL_NET",
            "direction": "OUTPUT",
            "signal_shape_info": {
                "shape": "scalar",
                "expected_physical_pin_count": 1,
                "parent_line_id": "L1",
                "member_index": 1,
                "member_count": 1,
                "is_expanded_member": False,
                "confidence": "auto_high",
                "needs_model_shape_review": False,
                "evidence": {"source_pins_available": True},
            },
        }
        compact = compact_normalized_connection(row)
        self.assertNotIn("base_line_id", compact)
        self.assertNotIn("output_sheet_name", compact)
        self.assertNotIn("raw_source_port", compact)
        self.assertNotIn("raw_target_port", compact)
        self.assertEqual(
            {"shape": "scalar", "expected_physical_pin_count": 1},
            compact["signal_shape_info"],
        )

        diagram = build_diagram_link_context(
            {
                "link_family_ids": ["CONTROL"],
                "link_family_sources": ["explicit_link_info"],
            },
            [row],
        )
        self.assertNotIn("line_link_contexts", diagram)

    def test_finish_checker_accepts_minimal_task_plan(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            output_file = root / "outputs" / "TASK_01.jsonl"
            output_file.parent.mkdir()
            output_file.write_text(
                json.dumps(
                    {
                        "line_id": "L1",
                        "selected_pin": "",
                        "decision_type": "unresolved",
                        "confidence": "Low",
                        "analysis": "insufficient information",
                        "needs_human_review": True,
                    }
                )
                + "\n",
                encoding="utf-8",
            )
            plan = root / "model_resolution_tasks" / "subagent_task_plan.json"
            plan.parent.mkdir()
            plan.write_text(
                json.dumps(
                    {
                        "tasks": [
                            {
                                "task_id": "TASK_01",
                                "line_ids": ["L1"],
                                "output_file": str(output_file),
                            }
                        ]
                    }
                ),
                encoding="utf-8",
            )
            report = check_subagent_outputs(
                plan,
                root / "model_resolved_decisions.jsonl",
                root / "subagent_output_check.json",
                root / "failed_subagent_rerun_plan.md",
            )
            self.assertEqual("PASS", report["status"])

    def test_finish_checker_blocks_missing_session_status(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            output_file = root / "outputs" / "TASK_01.jsonl"
            output_file.parent.mkdir()
            output_file.write_text(json.dumps({"line_id": "L1"}) + "\n", encoding="utf-8")
            plan = root / "model_resolution_tasks" / "subagent_task_plan.json"
            plan.parent.mkdir()
            plan.write_text(
                json.dumps(
                    {
                        "tasks": [
                            {
                                "task_id": "TASK_01",
                                "subagent_session_id": "SESSION_A",
                                "line_ids": ["L1"],
                                "output_file": str(output_file),
                            }
                        ]
                    }
                ),
                encoding="utf-8",
            )

            report = check_subagent_outputs(
                plan,
                root / "model_resolved_decisions.jsonl",
                root / "subagent_output_check.json",
                root / "failed_subagent_rerun_plan.md",
            )

            self.assertEqual("FAIL", report["status"])
            self.assertEqual("FAIL", report["session_status"]["status"])
            self.assertIn("missing subagent session status file", report["session_status"]["errors"][0])

    def test_finish_checker_requires_completed_session_status(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            output_file = root / "outputs" / "TASK_01.jsonl"
            output_file.parent.mkdir()
            output_file.write_text(json.dumps({"line_id": "L1"}) + "\n", encoding="utf-8")
            plan_dir = root / "model_resolution_tasks"
            plan_dir.mkdir()
            status_file = plan_dir / "subagent_session_status.json"
            plan = plan_dir / "subagent_task_plan.json"
            plan.write_text(
                json.dumps(
                    {
                        "subagent_session_status_json": str(status_file),
                        "tasks": [
                            {
                                "task_id": "TASK_01",
                                "subagent_session_id": "SESSION_A",
                                "line_ids": ["L1"],
                                "output_file": str(output_file),
                            }
                        ],
                    }
                ),
                encoding="utf-8",
            )
            status_file.write_text(
                json.dumps(
                    {
                        "sessions": [
                            {
                                "subagent_session_id": "SESSION_A",
                                "spawned": True,
                                "completed": True,
                                "failed": False,
                                "timed_out": False,
                                "rerun_required": False,
                                "completed_task_ids": ["TASK_01"],
                            }
                        ]
                    }
                ),
                encoding="utf-8",
            )

            report = check_subagent_outputs(
                plan,
                root / "model_resolved_decisions.jsonl",
                root / "subagent_output_check.json",
                root / "failed_subagent_rerun_plan.md",
            )

            self.assertEqual("PASS", report["status"])
            self.assertEqual("PASS", report["session_status"]["status"])

    def test_large_pin_group_catalog_is_lossless_multilabel_and_unranked(self) -> None:
        pins = ["RFIN0_P", "RFIN0_N", "SPI1_CLK", "VDD_1V8", "SPECIAL_X"]
        catalog = build_pin_group_catalog("PART_A", pins)

        self.assertEqual(PIN_GROUP_THRESHOLD, catalog["grouping_threshold"])
        self.assertTrue(catalog["lossless"])
        self.assertTrue(catalog["multi_label"])
        self.assertFalse(catalog["ranked"])
        self.assertEqual(pins, catalog["all_pins"])
        self.assertIn("RFIN0_P", catalog["groups"]["ANALOG_RF"])
        self.assertIn("RFIN0_P", catalog["groups"]["DIFFERENTIAL"])
        self.assertEqual(["SPECIAL_X"], catalog["groups"]["UNCLASSIFIED"])

        grouped_pins = {
            pin
            for group_pins in catalog["groups"].values()
            for pin in group_pins
        }
        self.assertEqual(set(pins), grouped_pins)
        self.assertEqual(
            ["RFIN0_P", "RFIN0_N"],
            pins_for_groups(catalog, ["DIFFERENTIAL"]),
        )


    def test_pin_renderer_keeps_standard_network_column_empty_and_preserves_analysis(self) -> None:
        rows = rows_for_decision(
            {
                "source_block_name": "SROC",
                "source_port": "SPI1",
                "target_block_name": "AMC7964",
                "target_port": "SPI",
                "connection_id": "C1",
                "connection_attribute": "控制信号",
            },
            {
                "selected_pin": "HAC_SPI1_CLK",
                "analysis": "pin mapping basis",
                "confidence": "High",
            },
        )
        self.assertEqual("", rows[0]["网络命名"])
        self.assertEqual("pin mapping basis", rows[0]["分析说明"])
        self.assertEqual("控制信号", rows[0]["连线属性"])

    def test_internal_flat_headers_keep_canonical_fields_for_debug_outputs(self) -> None:
        self.assertEqual(14, len(FINAL_HEADERS))
        self.assertEqual(
            [
                "源Block标识", "源Block名称", "源Port", "目的Block标识", "目的Block名称",
                "目的Port", "连线ID", "连线名称", "连线方向", "连线属性",
            ],
            FINAL_HEADERS[:10],
        )
        self.assertEqual(["原理图Pin脚", "分析说明", "映射置信度", "网络命名"], FINAL_HEADERS[10:])

    def test_normalizer_reads_new_connection_attribute_and_keeps_old_input_compatible(self) -> None:
        source = self.task_dir / "connections.json"
        output = self.task_dir / "normalized.jsonl"
        source.write_text(json.dumps([
            {
                "source_sheet_name": "S1", "源Block标识": "B1", "源Block名称": "SRC",
                "源Port": "GPIO0", "目的Block标识": "B2", "目的Block名称": "DST",
                "目的Port": "EN", "连线ID": "C1", "连线名称": "ENABLE",
                "连线方向": "OUTPUT", "连线属性": "低有效",
            },
            {
                "source_sheet_name": "S1", "源Block标识": "B1", "源Block名称": "SRC",
                "源Port": "GPIO1", "目的Block标识": "B2", "目的Block名称": "DST",
                "目的Port": "MODE", "连线ID": "C2", "连线名称": "MODE",
                "连线方向": "OUTPUT",
            },
        ], ensure_ascii=False), encoding="utf-8")

        rows = normalize_connections(source, output)

        self.assertEqual("低有效", rows[0]["connection_attribute"])
        self.assertEqual("", rows[1]["connection_attribute"])

    def test_excel_reads_business_fields_by_header_name_when_columns_are_reordered(self) -> None:
        from openpyxl import Workbook

        source = self.task_dir / "reordered_connections.xlsx"
        output = self.task_dir / "normalized.jsonl"
        workbook = Workbook()
        block_info = workbook.active
        block_info.title = "block_info"
        block_info.append(["说明"])
        block_info.append(["器件信息", "无关列", "框图标识"])
        block_info.append(["PART_SRC", "ignored", "SRC_SHEET"])

        connections = workbook.create_sheet("SRC_SHEET")
        reordered_headers = [
            "连线属性", "目的Port", "源Block名称", "连线ID", "源Port",
            "目的Block名称", "连线方向", "源Block标识", "连线名称", "目的Block标识",
        ]
        connections.append(["本页说明"])
        connections.append(reordered_headers)
        values_by_header = {
            "源Block标识": "SRC", "源Block名称": "SOURCE", "源Port": "GPIO0",
            "目的Block标识": "DST", "目的Block名称": "TARGET", "目的Port": "EN",
            "连线ID": "C1", "连线名称": "ENABLE", "连线方向": "OUTPUT", "连线属性": "低有效",
        }
        connections.append([values_by_header[header] for header in reordered_headers])

        link_info = workbook.create_sheet("link_info")
        link_info.append(["本页说明"])
        link_headers = ["器件角色说明", "相关连线ID", "链路编号", "器件Sheet", "链路类型"]
        link_info.append(link_headers)
        link_values = {
            "链路类型": "CONTROL", "链路编号": "LINK_1", "器件Sheet": "SRC_SHEET",
            "相关连线ID": "C1", "器件角色说明": "source controller",
        }
        link_info.append([link_values[header] for header in link_headers])
        workbook.save(source)

        rows = normalize_connections(source, output)

        self.assertEqual(1, len(rows))
        self.assertEqual(3, rows[0]["source_row_number"])
        self.assertEqual("PART_SRC", rows[0]["source_part_id"])
        self.assertEqual("GPIO0", rows[0]["source_port"])
        self.assertEqual("EN", rows[0]["target_port"])
        self.assertEqual("低有效", rows[0]["connection_attribute"])
        self.assertEqual("CONTROL", rows[0]["link_family_id"])
        self.assertEqual("LINK_1", rows[0]["link_instance_id"])

        pins_source = self.task_dir / "reordered_pins.xlsx"
        pins_workbook = Workbook()
        pins_sheet = pins_workbook.active
        pins_sheet.title = "pins"
        pins_sheet.append(["说明"])
        pins_sheet.append(["pin_name", "componentUuid"])
        pins_sheet.append(["GPIO0", "PART_SRC"])
        pins_workbook.save(pins_source)
        self.assertEqual(["GPIO0"], load_pin_catalog(pins_source)["PART_SRC"])

    def test_incremental_excel_filters_delete_and_uses_existing_pin_as_context(self) -> None:
        from openpyxl import Workbook, load_workbook

        source = self.task_dir / "incremental.xlsx"
        normalized = self.task_dir / "incremental_normalized.jsonl"
        pre_decisions = self.task_dir / "incremental_pre_decisions.jsonl"
        needs_model = self.task_dir / "incremental_needs_model.jsonl"
        pins = self.task_dir / "incremental_pins.json"
        output = self.task_dir / "incremental_output.xlsx"
        final_decisions = self.task_dir / "incremental_final_decisions.jsonl"

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "DEV0"
        headers = [
            "源Port", "目的Port", "连线ID", "原理图pin脚", "分析说明",
            "映射置信度", "网络命名", "修改类型", "自定义字段",
        ]
        sheet.append(headers)
        sheet.append(["OLD_LOGICAL", "OLD_TARGET", "C_OLD", "PIN_USED", "已有结论", "High", "OLD_NET", "modify", "KEEP_OLD"])
        sheet.append(["NEW_LOGICAL", "NEW_TARGET", "C_NEW", "", "", "", "", "add", "KEEP_NEW"])
        sheet.append(["DELETED", "NO_TARGET", "C_DELETE", "", "", "", "", "delete", "KEEP_DELETE"])
        block_info = workbook.create_sheet("block_info")
        block_info.append(["框图标识", "器件信息"])
        block_info.append(["DEV0", "PART_A"])
        workbook.save(source)
        pins.write_text(json.dumps({"PART_A": ["PIN_USED", "PIN_NEW"]}), encoding="utf-8")

        rows = normalize_connections(source, normalized)
        route_model_resolution(normalized, pins, pre_decisions, needs_model)

        self.assertEqual(["C_OLD", "C_NEW"], [row["connection_id"] for row in rows])
        self.assertFalse(rows[0]["analysis_required"])
        self.assertEqual("PIN_USED", rows[0]["existing_selected_pin"])
        self.assertTrue(rows[1]["analysis_required"])
        routed = list(iter_jsonl(needs_model))
        self.assertEqual([rows[1]["line_id"]], [row["line_id"] for row in routed])
        decisions_by_id = {row["line_id"]: row for row in iter_jsonl(pre_decisions)}
        self.assertEqual("existing_mapping", decisions_by_id[rows[0]["line_id"]]["decision_type"])

        allocation = build_pin_allocation_context(rows)
        self.assertEqual("PIN_USED", allocation["existing_pin_assignments"][0]["selected_pin"])

        write_jsonl(final_decisions, [
            decisions_by_id[rows[0]["line_id"]],
            {
                "line_id": rows[1]["line_id"], "selected_pin": "PIN_NEW",
                "analysis": "增量分析", "confidence": "High",
            },
        ])
        render_template_sheets(source, normalized, final_decisions, output)

        rendered = load_workbook(output, data_only=True)
        try:
            rendered_sheet = rendered["DEV0"]
            header_map = worksheet_header_map(rendered_sheet, 1)
            pin_headers = [header for header in header_map if header.casefold() == "原理图pin脚".casefold()]
            self.assertEqual(["原理图pin脚"], pin_headers)
            pin_column = header_map[pin_headers[0]]
            self.assertEqual("PIN_USED", rendered_sheet.cell(row=2, column=pin_column).value)
            self.assertEqual("PIN_NEW", rendered_sheet.cell(row=3, column=pin_column).value)
            self.assertIsNone(rendered_sheet.cell(row=4, column=pin_column).value)
            self.assertEqual("KEEP_DELETE", rendered_sheet.cell(row=4, column=header_map["自定义字段"]).value)
            self.assertEqual("OLD_NET", rendered_sheet.cell(row=2, column=header_map["网络命名"]).value)
        finally:
            rendered.close()

    def test_formal_excel_renderer_preserves_input_columns_and_appends_results(self) -> None:
        from openpyxl import Workbook, load_workbook

        template = self.task_dir / "template.xlsx"
        normalized = self.task_dir / "normalized.jsonl"
        decisions = self.task_dir / "decisions.jsonl"
        result = self.task_dir / "result.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "S1"
        input_headers = ["自定义字段", "连线ID", "分析说明", "源Port", "目的Port", "连线属性", "备注"]
        sheet.append(["输入说明"])
        sheet.append(input_headers)
        input_values = {
            "自定义字段": "KEEP_ME", "连线ID": "C1", "分析说明": "旧结果", "源Port": "GPIO0",
            "目的Port": "EN", "连线属性": "低有效", "备注": "原始备注",
        }
        sheet.append([input_values[header] for header in input_headers])
        workbook.save(template)
        write_jsonl(normalized, [{
            "line_id": "S1:2:C1", "source_sheet_name": "S1", "output_sheet_name": "S1",
            "source_row_number": 3,
            "source_block_id": "B1", "source_block_name": "SRC", "source_port": "GPIO0",
            "target_block_id": "B2", "target_block_name": "DST", "target_port": "EN",
            "connection_id": "C1", "connection_name": "ENABLE", "direction": "OUTPUT",
            "connection_attribute": "低有效",
        }])
        write_jsonl(decisions, [{
            "line_id": "S1:2:C1", "selected_pin": "GPIO0", "analysis": "逐字命中",
            "confidence": "High", "decision_type": "model_resolved",
        }])

        render_template_sheets(template, normalized, decisions, result)

        rendered = load_workbook(result, data_only=True)
        try:
            sheet = rendered["S1"]
            header_map = worksheet_header_map(sheet, 2)
            self.assertEqual(input_headers, list(header_map)[:len(input_headers)])
            self.assertEqual(["原理图Pin脚", "映射置信度", "网络命名"], list(header_map)[len(input_headers):])
            self.assertEqual("输入说明", sheet.cell(row=1, column=1).value)
            self.assertEqual("KEEP_ME", sheet.cell(row=3, column=header_map["自定义字段"]).value)
            self.assertEqual("原始备注", sheet.cell(row=3, column=header_map["备注"]).value)
            self.assertEqual("低有效", sheet.cell(row=3, column=header_map["连线属性"]).value)
            self.assertEqual("GPIO0", sheet.cell(row=3, column=header_map["原理图Pin脚"]).value)
            self.assertEqual("逐字命中", sheet.cell(row=3, column=header_map["分析说明"]).value)
            self.assertEqual("High", sheet.cell(row=3, column=header_map["映射置信度"]).value)
            self.assertIsNone(sheet.cell(row=3, column=header_map["网络命名"]).value)
        finally:
            rendered.close()

    def test_formal_excel_renderer_copies_complete_input_row_for_multi_pin_expansion(self) -> None:
        from openpyxl import Workbook, load_workbook

        template = self.task_dir / "expanded_template.xlsx"
        normalized = self.task_dir / "expanded_normalized.jsonl"
        decisions = self.task_dir / "expanded_decisions.jsonl"
        result = self.task_dir / "expanded_result.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "S1"
        sheet.append(["备注", "源Port", "连线ID", "目的Port"])
        sheet.append(["must survive", "BUS", "C1", "TARGET"])
        workbook.save(template)
        write_jsonl(normalized, [{
            "line_id": "S1:1:C1", "source_sheet_name": "S1", "output_sheet_name": "S1",
            "source_row_number": 2, "source_port": "BUS", "target_port": "TARGET",
            "connection_id": "C1", "base_connection_id": "C1",
        }])
        write_jsonl(decisions, [{
            "line_id": "S1:1:C1", "selected_pins": ["PIN_A", "PIN_B"],
            "analyses": ["first", "second"], "confidences": ["High", "Medium"],
        }])

        render_template_sheets(template, normalized, decisions, result)

        rendered = load_workbook(result, data_only=True)
        try:
            sheet = rendered["S1"]
            header_map = worksheet_header_map(sheet, 1)
            self.assertEqual("must survive", sheet.cell(row=2, column=header_map["备注"]).value)
            self.assertEqual("must survive", sheet.cell(row=3, column=header_map["备注"]).value)
            self.assertEqual("C1#1", sheet.cell(row=2, column=header_map["连线ID"]).value)
            self.assertEqual("C1#2", sheet.cell(row=3, column=header_map["连线ID"]).value)
            self.assertEqual("PIN_A", sheet.cell(row=2, column=header_map["原理图Pin脚"]).value)
            self.assertEqual("PIN_B", sheet.cell(row=3, column=header_map["原理图Pin脚"]).value)
        finally:
            rendered.close()

    def test_renderer_groups_preexpanded_members_from_one_input_row(self) -> None:
        normalized = self.task_dir / "preexpanded_normalized.jsonl"
        decisions = self.task_dir / "preexpanded_decisions.jsonl"
        base = {
            "source_sheet_name": "S1", "output_sheet_name": "S1", "source_row_number": 2,
            "parent_line_id": "S1:1:C1", "base_line_id": "S1:1:C1", "base_connection_id": "C1",
            "source_port": "BUS", "target_port": "TARGET",
        }
        write_jsonl(normalized, [
            {**base, "line_id": "S1:1:C1#1", "connection_id": "C1#1"},
            {**base, "line_id": "S1:1:C1#2", "connection_id": "C1#2"},
        ])
        write_jsonl(decisions, [
            {"line_id": "S1:1:C1#1", "selected_pin": "PIN_A", "analysis": "first", "confidence": "High"},
            {"line_id": "S1:1:C1#2", "selected_pin": "PIN_B", "analysis": "second", "confidence": "High"},
        ])

        grouped = build_final_rows_by_template_sheet(normalized, decisions)

        self.assertEqual(1, len(grouped["S1"]))
        self.assertEqual(["PIN_A", "PIN_B"], [row["原理图Pin脚"] for row in grouped["S1"][0]["result_rows"]])

    def test_compact_task_connection_preserves_connection_attribute(self) -> None:
        compact = compact_normalized_connection({
            "line_id": "S:2:C1",
            "source_sheet_name": "S",
            "source_part_id": "P1",
            "source_block_id": "B1",
            "source_block_name": "SRC",
            "source_port": "GPIO0",
            "target_block_id": "B2",
            "target_block_name": "DST",
            "target_port": "EN",
            "connection_id": "C1",
            "connection_name": "ENABLE",
            "direction": "OUTPUT",
            "connection_attribute": "低有效",
        })
        self.assertEqual("低有效", compact["connection_attribute"])

    def test_model_tasks_persist_canonical_finish_contract(self) -> None:
        connections = self.task_dir / "design" / "aggregated_connections.xlsx"
        pins = self.task_dir / "pin_info.json"
        config_path = save_pipeline_run_config(self.task_dir, str(connections), str(pins))

        self.assertEqual(self.task_dir / "intermediate" / PIPELINE_CONFIG_FILENAME, config_path)
        config = load_pipeline_run_config(self.task_dir)
        self.assertEqual(str(connections.resolve()), config["connections"])
        self.assertEqual(str(connections.resolve()), config["template_excel"])
        self.assertEqual(str((ROOT / "scripts" / "run_pipeline.py").resolve()), config["script_path"])
        command = (self.task_dir / "intermediate" / FINISH_COMMAND_FILENAME).read_text(encoding="utf-8")
        self.assertNotIn("Set-Location", command)
        self.assertIn(str((ROOT / "scripts" / "run_pipeline.py").resolve()), command)
        self.assertIn("--stage finish", command)
        self.assertNotIn("--normalized", command)
        self.assertNotIn("--decisions", command)

    def test_finish_never_renders_when_validation_fails(self) -> None:
        with (
            mock.patch("run_pipeline.run_apply"),
            mock.patch("run_pipeline.merge_decisions"),
            mock.patch("run_pipeline.validate_mapping", return_value={"status": "ERROR"}),
            mock.patch("run_pipeline.render_template_sheets") as render,
        ):
            with self.assertRaisesRegex(RuntimeError, "validation_report.json is not PASS"):
                run_finish(
                    self.task_dir,
                    "connections.xlsx",
                    "pin_info.json",
                    "connections.xlsx",
                    "template_sheets",
                )
            render.assert_not_called()


if __name__ == "__main__":
    unittest.main()
