#!/usr/bin/env python3
# -*- coding: utf-8 -*-
from __future__ import annotations
import csv,json,re,hashlib
from pathlib import Path
from typing import Any,Dict,Iterable,List,Optional

FINAL_HEADERS=["源Block标识","源Block名称","源Port","目的Block标识","目的Block名称","目的Port","连线ID","连线名称","连线方向","连线属性","原理图Pin脚","分析说明","映射置信度","网络命名"]
VALID_DIRECTIONS={"INPUT","OUTPUT",""}; VALID_CONFIDENCE={"High","Medium","Low",""}
EXCLUDED_INPUT_SHEETS={"BLOCK_INFO","目录","说明","README","INDEX"}

def ensure_dir(path): p=Path(path); p.mkdir(parents=True,exist_ok=True); return p

def read_json(path,default=None):
    p=Path(path); return default if not p.exists() else json.loads(p.read_text(encoding='utf-8'))

def write_json(path,obj): p=Path(path); ensure_dir(p.parent); p.write_text(json.dumps(obj,ensure_ascii=False,indent=2),encoding='utf-8')

def iter_jsonl(path):
    p=Path(path)
    if not p.exists(): return
    with p.open('r',encoding='utf-8') as f:
        for line in f:
            line=line.strip()
            if line: yield json.loads(line)

def write_jsonl(path,rows):
    p=Path(path); ensure_dir(p.parent)
    with p.open('w',encoding='utf-8') as f:
        for row in rows: f.write(json.dumps(row,ensure_ascii=False)+'\n')

def stable_hash(obj): return hashlib.md5(json.dumps(obj,ensure_ascii=False,sort_keys=True).encode()).hexdigest()[:12]
def normalize_text(v): return '' if v is None else str(v).strip()

COMPONENT_UUID_KEYS={"componentuuid","component_uuid","component-uuid"}
PIN_COLLECTION_KEYS=("pins","pin_info","pinInfo","pin_list","pinList","pinNames","pin_names")
PIN_NAME_KEYS=("pin","pin_name","pinName","Pin","name")
COMPONENT_CONTAINER_KEYS={"components","componentpins","component_pin_infos","devices","items","data","records"}

def _clean_component_uuid(value):
    text=normalize_text(value).strip('\\').strip().strip('"\'').strip()
    if len(text)>=2 and text[0]=='{' and text[-1]=='}': text=text[1:-1].strip()
    return text

def _component_uuid_from_object(value):
    if isinstance(value,dict):
        for key,item in value.items():
            if normalize_text(key).lower() in COMPONENT_UUID_KEYS:
                found=_clean_component_uuid(item)
                if found: return found
        for item in value.values():
            found=_component_uuid_from_object(item)
            if found: return found
    elif isinstance(value,list):
        for item in value:
            found=_component_uuid_from_object(item)
            if found: return found
    return ''

def extract_component_uuid(value):
    """Extract componentUuid from JSON-like block_info device metadata."""
    if isinstance(value,(dict,list)):
        return _component_uuid_from_object(value)
    text=normalize_text(value)
    if not text or 'componentuuid' not in text.lower().replace('_','').replace('-',''):
        return ''
    decoded=text
    for _ in range(2):
        try:
            parsed=json.loads(decoded)
        except (json.JSONDecodeError,TypeError):
            break
        found=_component_uuid_from_object(parsed)
        if found: return found
        if not isinstance(parsed,str): break
        decoded=parsed
    quoted=re.search(r'''(?ix)
        [\\"']*component[_-]?uuid[\\"']*\s*[:=]\s*
        (?:\\?["'])(?P<value>[^"'\\]+)(?:\\?["'])
    ''',text)
    if quoted:
        return _clean_component_uuid(quoted.group('value'))
    bare=re.search(r'''(?ix)
        \bcomponent[_-]?uuid\b\s*[:=]\s*
        (?P<value>\{?[a-z0-9][a-z0-9._:/-]*\}?)
    ''',text)
    return _clean_component_uuid(bare.group('value')) if bare else ''

def normalize_component_reference(value):
    return extract_component_uuid(value) or normalize_text(value)
def normalize_direction(v):
    s=normalize_text(v).upper()
    if s in {'IN','INPUT'}: return 'INPUT'
    if s in {'OUT','OUTPUT'}: return 'OUTPUT'
    return s if s in VALID_DIRECTIONS else ''

def _read_xlsx_rows(path: Path) -> List[Dict[str, Any]]:
    """读取输入框图 Excel：保留分页信息。默认跳过 BLOCK_INFO 等非器件页。"""
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    rows: List[Dict[str, Any]] = []
    for ws in wb.worksheets:
        if ws.title.strip().upper() in {x.upper() for x in EXCLUDED_INPUT_SHEETS}:
            continue
        values = list(ws.iter_rows(values_only=True))
        if not values:
            continue
        header_idx = None
        headers = None
        for i, row in enumerate(values[:20]):
            cells=[normalize_text(c) for c in row]
            if any(c in {"源Port","source_port","目的Port","目标Port","target_port","连线ID","连线标识","connection_id"} for c in cells):
                header_idx=i; headers=cells; break
        if header_idx is None or not headers:
            continue
        for raw in values[header_idx+1:]:
            if raw is None:
                continue
            row_dict={headers[i]: raw[i] if i < len(raw) else '' for i in range(len(headers)) if headers[i]}
            if not any(normalize_text(v) for v in row_dict.values()):
                continue
            row_dict['__sheet_name']=ws.title
            row_dict.setdefault('source_sheet_name', ws.title)
            row_dict.setdefault('output_sheet_name', ws.title)
            rows.append(row_dict)
    wb.close()
    return rows

def read_table(path):
    p=Path(path); suf=p.suffix.lower()
    if suf=='.json':
        data=json.loads(p.read_text(encoding='utf-8'))
        if isinstance(data,list): return data
        if isinstance(data,dict):
            for k in ('rows','connections','pins'):
                if isinstance(data.get(k),list): return data[k]
        raise ValueError(f'Unsupported JSON shape: {p}')
    if suf=='.jsonl': return list(iter_jsonl(p))
    if suf=='.csv':
        with p.open('r',encoding='utf-8-sig',newline='') as f: return list(csv.DictReader(f))
    if suf in {'.xlsx','.xlsm'}:
        return _read_xlsx_rows(p)
    raise ValueError(f'Unsupported file type: {p.suffix}; use json/jsonl/csv/xlsx.')

def _pin_names_from_value(value):
    if isinstance(value,list):
        result=[]
        for item in value:
            if isinstance(item,dict):
                pin=next((normalize_text(item.get(key)) for key in PIN_NAME_KEYS if normalize_text(item.get(key))), '')
                if pin: result.append(pin)
            else:
                pin=normalize_text(item)
                if pin: result.append(pin)
        return result
    if isinstance(value,dict):
        direct=next((normalize_text(value.get(key)) for key in PIN_NAME_KEYS if normalize_text(value.get(key))), '')
        if direct: return [direct]
        for key in PIN_COLLECTION_KEYS:
            if key in value:
                nested=_pin_names_from_value(value.get(key))
                if nested: return nested
    return []

def _collect_component_pin_records(value,catalog):
    if isinstance(value,list):
        for item in value: _collect_component_pin_records(item,catalog)
        return
    if not isinstance(value,dict): return
    component_id=''
    for key,item in value.items():
        if normalize_text(key).lower() in COMPONENT_UUID_KEYS:
            component_id=_clean_component_uuid(item)
            break
    pin_values=[]
    for key in PIN_COLLECTION_KEYS:
        if key in value:
            pin_values=_pin_names_from_value(value.get(key))
            if pin_values: break
    if not pin_values:
        direct_pin=next((normalize_text(value.get(key)) for key in PIN_NAME_KEYS if normalize_text(value.get(key))), '')
        if direct_pin: pin_values=[direct_pin]
    if component_id and pin_values:
        catalog.setdefault(component_id,[]).extend(pin_values)
    for item in value.values():
        if isinstance(item,(dict,list)): _collect_component_pin_records(item,catalog)

def load_pin_catalog(path):
    """Return pin_info as the native {part_number: [pin_name]} catalog."""
    p=Path(path); suf=p.suffix.lower()
    if suf=='.json':
        data=json.loads(p.read_text(encoding='utf-8'))
        if (
            isinstance(data,dict)
            and not any(normalize_text(key).lower() in COMPONENT_CONTAINER_KEYS for key in data)
            and all(isinstance(v,list) and all(not isinstance(item,(dict,list)) for item in v) for v in data.values())
        ):
            return {normalize_component_reference(k):[normalize_text(x) for x in v if normalize_text(x)] for k,v in data.items()}
        if isinstance(data,dict) and isinstance(data.get('pins'),dict):
            return {normalize_component_reference(k):[normalize_text(x) for x in v if normalize_text(x)] for k,v in data['pins'].items()}
        catalog={}
        _collect_component_pin_records(data,catalog)
        if catalog: return catalog
    catalog={}
    for row in read_table(p):
        part=normalize_component_reference(row.get('componentUuid') or row.get('component_uuid') or row.get('device_part_id') or row.get('part_number') or row.get('器件信息') or row.get('device_id') or row.get('block_id') or row.get('source_block_id') or row.get('器件标识'))
        pin=normalize_text(row.get('pin') or row.get('pin_name') or row.get('pinName') or row.get('Pin') or row.get('name'))
        if part and pin: catalog.setdefault(part,[]).append(pin)
    return catalog

def resolve_catalog_key(catalog, part_id):
    part=normalize_component_reference(part_id)
    if part in catalog: return part
    part_lower=part.lower()
    for key in catalog:
        if normalize_text(key).lower()==part_lower: return key
    stripped=part.lstrip('0')
    for key in catalog:
        if key.lstrip('0')==stripped:
            return key
    return part

def pins_for_part(catalog, part_id):
    return catalog.get(resolve_catalog_key(catalog, part_id), [])
