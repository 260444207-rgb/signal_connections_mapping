#!/usr/bin/env python3
# -*- coding: utf-8 -*-

from __future__ import annotations

import argparse
import re
from pathlib import Path
from typing import Dict, List, Any, Iterable

from .common import write_jsonl, normalize_text, normalize_direction, normalize_component_reference, read_table

FIELD_ALIASES = {
    "source_block_id": ["source_block_id", "源Block标识", "起止Block标识", "起始Block标识", "source_id", "block_id", "框图标识"],
    "source_block_name": ["source_block_name", "源Block名称", "起始Block名称", "source_name", "器件名称"],
    "source_port": ["source_port", "源Port", "源端Port", "port"],
    "target_block_id": ["target_block_id", "目的Block标识", "目标Block标识", "target_id"],
    "target_block_name": ["target_block_name", "目的Block名称", "目标Block名称", "target_name"],
    "target_port": ["target_port", "目的Port", "目标Port"],
    "connection_id": ["connection_id", "连线ID", "连线标识", "line_id", "edge_id"],
    "connection_name": ["connection_name", "连线名称"],
    "direction": ["direction", "连线方向"],
    "connection_attribute": ["connection_attribute", "连线属性", "连接属性", "line_attribute"],
}

SKIP_SHEETS = {"BLOCK_INFO", "LINK_INFO", "链路信息", "说明", "README", "INDEX", "目录"}

LINK_INFO_SHEET_NAMES = {"link_info", "链路信息"}

LINK_INFO_ALIASES = {
    "link_family_id": ["link_family_id", "link_type", "链路类型", "链路族", "链路族ID"],
    "link_instance_id": ["link_instance_id", "link_id", "链路编号", "链路ID", "链路标识"],
    "member_sheets": ["member_sheets", "sheet_names", "成员Sheet", "相关Sheet", "器件Sheet", "器件Sheet名字", "相关器件Sheet"],
    "member_connection_ids": ["member_connection_ids", "connection_ids", "成员连线ID", "相关连线ID", "连线ID"],
    "user_link_info": ["user_link_info", "link_description", "description", "用户标识的链路信息", "链路信息", "用户链路说明", "说明"],
    "device_role_info": ["device_role_info", "role_info", "器件角色说明", "器件角色", "角色说明"],
}

def pick(row: Dict[str, Any], canonical: str) -> str:
    for k in FIELD_ALIASES[canonical]:
        if k in row and row[k] is not None:
            return normalize_text(row[k])
    return ""

def pick_alias(row: Dict[str, Any], aliases: list[str]) -> str:
    for key in aliases:
        if key in row and row[key] is not None:
            return normalize_text(row[key])
    return ""

def split_tokens(value: Any) -> list[str]:
    text = normalize_text(value)
    if not text:
        return []
    return [x.strip() for x in re.split(r"[\n,，;；、]+", text) if x.strip()]

def normalize_meta_key(value: Any) -> str:
    return normalize_text(value).lower()

def expansion_count(port: str) -> int:
    port = normalize_text(port)
    m = re.match(r"^(.+)\[(\d+):(\d+)\]$", port)
    if m:
        a, b = int(m.group(2)), int(m.group(3))
        return abs(a - b) + 1
    m = re.match(r"^(.+)\*(\d+)$", port)
    if m:
        return int(m.group(2))
    return 1

def read_block_info(path: str | Path) -> Dict[str, str]:
    p = Path(path)
    if p.suffix.lower() not in {".xlsx", ".xlsm"}:
        return {}
    from openpyxl import load_workbook
    wb = load_workbook(p, data_only=True)
    try:
        ws = None
        for candidate in wb.worksheets:
            if candidate.title.strip().lower() == "block_info":
                ws = candidate
                break
        if ws is None:
            return {}
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return {}
        headers = [normalize_text(v) for v in rows[0]]
        try:
            key_idx = headers.index("框图标识")
            part_idx = headers.index("器件信息")
        except ValueError:
            return {}
        result = {}
        for raw in rows[1:]:
            key = normalize_text(raw[key_idx] if key_idx < len(raw) else "")
            part = normalize_component_reference(raw[part_idx] if part_idx < len(raw) else "")
            if key and part:
                result[key.lower()] = part
        return result
    finally:
        wb.close()

def read_link_info(input_path: str | Path) -> Dict[tuple[str, str], List[Dict[str, Any]]]:
    """
    读取可选的 link_info / 链路信息 sheet。

    推荐表头：
    链路类型 | 链路编号 | 器件Sheet | 用户标识的链路信息 | 器件角色说明 | 相关连线ID

    - 链路类型相同表示同一个 link_family。
    - 链路编号表示一条具体链路实例。
    - 器件Sheet 可用逗号/分号/顿号分隔。
    - 相关连线ID 可写裸 ID，也可写 sheet:连线ID；为空时表示这些 sheet 参与该链路，但具体连接归属由 subagent 结合角色说明判断。
    """
    p = Path(input_path)
    if p.suffix.lower() not in {".xlsx", ".xlsm"}:
        return {}

    from openpyxl import load_workbook

    wb = load_workbook(p, data_only=True)
    try:
        ws = None
        for candidate in wb.worksheets:
            if candidate.title.strip().lower() in LINK_INFO_SHEET_NAMES:
                ws = candidate
                break
        if ws is None:
            return {}

        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return {}

        header_idx = None
        headers = None
        for i, raw in enumerate(rows):
            values = [normalize_text(v) for v in raw]
            if any(values):
                header_idx = i
                headers = values
                break
        if header_idx is None or not headers:
            return {}

        result: Dict[tuple[str, str], List[Dict[str, Any]]] = {}
        for raw in rows[header_idx + 1:]:
            values = [normalize_text(v) for v in raw]
            if not any(values):
                continue
            row = {h: v for h, v in zip(headers, values) if h}

            link_family_id = pick_alias(row, LINK_INFO_ALIASES["link_family_id"])
            link_instance_id = pick_alias(row, LINK_INFO_ALIASES["link_instance_id"])
            member_sheets = split_tokens(pick_alias(row, LINK_INFO_ALIASES["member_sheets"]))
            member_connection_ids = split_tokens(pick_alias(row, LINK_INFO_ALIASES["member_connection_ids"]))
            user_link_info = pick_alias(row, LINK_INFO_ALIASES["user_link_info"])
            device_role_info = pick_alias(row, LINK_INFO_ALIASES["device_role_info"])

            if not link_family_id and not link_instance_id:
                continue

            meta = {
                "link_family_id": link_family_id,
                "link_instance_id": link_instance_id,
                "link_member_sheets": member_sheets,
                "user_link_info": user_link_info,
                "device_role_info": device_role_info,
            }

            if not member_connection_ids:
                for sheet in member_sheets:
                    result.setdefault((normalize_meta_key(sheet), ""), []).append(meta)
                continue

            for item in member_connection_ids:
                if ":" in item:
                    sheet, connection_id = item.split(":", 1)
                    result.setdefault((normalize_meta_key(sheet), normalize_meta_key(connection_id)), []).append(meta)
                else:
                    for sheet in member_sheets:
                        result.setdefault((normalize_meta_key(sheet), normalize_meta_key(item)), []).append(meta)
        return result
    finally:
        wb.close()

def lookup_link_meta(
    link_info: Dict[tuple[str, str], List[Dict[str, Any]]],
    sheet_name: str,
    connection_id: str,
) -> Dict[str, Any]:
    sheet_key = normalize_meta_key(sheet_name)
    connection_key = normalize_meta_key(connection_id)
    exact_contexts = link_info.get((sheet_key, connection_key), [])
    sheet_contexts = link_info.get((sheet_key, ""), [])
    contexts = exact_contexts or sheet_contexts
    if not contexts:
        return {}
    if len(contexts) == 1:
        meta = dict(contexts[0])
        meta["link_contexts"] = contexts
        return meta
    return {
        "link_family_id": "",
        "link_instance_id": "",
        "link_member_sheets": sorted({sheet for ctx in contexts for sheet in ctx.get("link_member_sheets", [])}),
        "user_link_info": "",
        "device_role_info": "",
        "link_contexts": contexts,
    }

def lookup_block_part(block_parts: Dict[str, str], *keys: Any) -> str:
    """
    框图标识 / block 名称只用于在 block_info 中查出器件信息。
    后续器件类型判断使用查出的器件信息，不再从框图标识本身推断。
    """
    for key in keys:
        normalized = normalize_text(key).lower()
        if normalized and normalized in block_parts:
            return block_parts[normalized]
    for key in keys:
        normalized = normalize_text(key).lower()
        if not normalized:
            continue
        compact = re.sub(r"[^a-z0-9]+", "", normalized)
        if not re.search(r"[a-z]", compact):
            continue
        for known_key, part in block_parts.items():
            known_compact = re.sub(r"[^a-z0-9]+", "", known_key.lower())
            if not re.search(r"[a-z]", known_compact):
                continue
            if known_compact and (known_compact in compact or compact in known_compact):
                return part
    return ""

def add_block_part_alias(block_parts: Dict[str, str], key: Any, part: str) -> None:
    normalized = normalize_text(key).lower()
    if normalized and part:
        block_parts.setdefault(normalized, part)

def build_block_part_aliases(
    sheet_parts: Dict[str, str],
    sheet_rows: list[tuple[str, list[dict[str, Any]]]],
) -> Dict[str, str]:
    """
    block_info 给出的是 框图标识 -> 器件信息。
    连接行里真正出现的是 block_id/block_name，因此先用每个 sheet 的器件信息
    反向登记该 sheet 内源端 block_id/block_name 的别名。
    """
    aliases = dict(sheet_parts)
    for sheet_name, rows in sheet_rows:
        part = lookup_block_part(sheet_parts, sheet_name)
        if not part:
            continue
        add_block_part_alias(aliases, sheet_name, part)
        for row in rows:
            add_block_part_alias(aliases, pick(row, "source_block_id"), part)
            add_block_part_alias(aliases, pick(row, "source_block_name"), part)
    return aliases

def header_field_hits(cells: list[str]) -> set[str]:
    hits: set[str] = set()
    normalized_cells = {normalize_text(cell) for cell in cells if normalize_text(cell)}
    for field, aliases in FIELD_ALIASES.items():
        if any(alias in normalized_cells for alias in aliases):
            hits.add(field)
    return hits

def looks_like_connection_header(cells: list[str]) -> bool:
    hits = header_field_hits(cells)
    if "source_port" not in hits:
        return False
    endpoint_hits = {
        "source_block_id",
        "source_block_name",
        "target_block_id",
        "target_block_name",
        "target_port",
    }
    return "connection_id" in hits or bool(endpoint_hits & hits)

def read_excel_sheets(path: str | Path) -> Iterable[tuple[str, list[dict[str, Any]]]]:
    """
    读取输入框图 Excel 的原始 sheet。
    关键原则：输出分页只以这里读取到的 sheet_name 为准，不根据 block_id/group 自动造 sheet。
    """
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    for sheet_name in wb.sheetnames:
        if sheet_name.upper() in SKIP_SHEETS:
            continue
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue

        header_idx = None
        headers: list[str] | None = None
        first_non_empty_idx = None
        first_non_empty_headers: list[str] | None = None
        for i, r in enumerate(rows[:30]):
            values = [normalize_text(v) for v in r]
            if any(values) and first_non_empty_idx is None:
                first_non_empty_idx = i
                first_non_empty_headers = values
            if looks_like_connection_header(values):
                header_idx = i
                headers = values
                break
        if header_idx is None:
            header_idx = first_non_empty_idx
            headers = first_non_empty_headers
        if header_idx is None or not headers:
            continue

        data_rows = []
        for r in rows[header_idx + 1:]:
            values = [normalize_text(v) for v in r]
            if not any(values):
                continue
            row = {}
            for h, v in zip(headers, values):
                if h:
                    row[h] = v
            data_rows.append(row)

        yield sheet_name, data_rows

def read_connections_with_sheet(input_path: str | Path) -> Iterable[tuple[str, list[dict[str, Any]]]]:
    p = Path(input_path)
    if p.suffix.lower() in {".xlsx", ".xlsm"}:
        yield from read_excel_sheets(p)
    else:
        # JSON/CSV 也必须显式给 source_sheet_name/output_sheet_name；没有则只作为 DEFAULT_INPUT。
        rows = read_table(p)
        grouped: Dict[str, list[dict[str, Any]]] = {}
        for row in rows:
            sheet = normalize_text(row.get("source_sheet_name") or row.get("output_sheet_name") or row.get("sheet_name") or "DEFAULT_INPUT")
            grouped.setdefault(sheet, []).append(row)
        for sheet, sheet_rows in grouped.items():
            yield sheet, sheet_rows

def normalize_connections(input_path: str | Path, output_path: str | Path) -> List[Dict[str, Any]]:
    out = []
    sheet_parts = read_block_info(input_path)
    link_info = read_link_info(input_path)
    sheet_rows = list(read_connections_with_sheet(input_path))
    block_part_aliases = build_block_part_aliases(sheet_parts, sheet_rows)

    for sheet_name, rows in sheet_rows:
        for idx, row in enumerate(rows, start=1):
            source_port = pick(row, "source_port")
            if not source_port:
                # 当前 sheet 中非连接行跳过，但不影响最终输出保留该 sheet。
                continue

            target_port_raw = pick(row, "target_port")
            source_block_id = pick(row, "source_block_id") or sheet_name
            source_block_name = pick(row, "source_block_name")
            target_block_id = pick(row, "target_block_id")
            target_block_name = pick(row, "target_block_name")
            source_part_id = lookup_block_part(block_part_aliases, source_block_id, source_block_name, sheet_name)
            target_part_id = lookup_block_part(block_part_aliases, target_block_id, target_block_name)
            # 显式范围只声明候选宽度；最终成员数由严格器件规则与 pin_info 共同确认。
            declared_bus_width = expansion_count(source_port)
            count = 1

            for bit_idx in range(1, count + 1):
                connection_id_raw = pick(row, "connection_id") or f"{sheet_name}_{idx}"
                display_connection_id = connection_id_raw if count == 1 else f"{connection_id_raw}#{bit_idx}"
                link_meta = lookup_link_meta(link_info, sheet_name, connection_id_raw)

                normalized = {
                    "line_id": f"{sheet_name}:{idx}:{display_connection_id}",
                    "base_line_id": f"{sheet_name}:{idx}:{connection_id_raw}",
                    "expansion_index": bit_idx,
                    "expansion_count": count,
                    "declared_bus_width": declared_bus_width if declared_bus_width > 1 else 0,
                    "analysis_unit": "device",
                    "source_sheet_name": sheet_name,
                    "output_sheet_name": sheet_name,
                    "source_part_id": source_part_id,
                    "target_part_id": target_part_id,
                    "source_block_id": source_block_id,
                    "source_block_name": source_block_name,
                    "source_port": source_port,
                    "target_block_id": target_block_id,
                    "target_block_name": target_block_name,
                    "target_port": target_port_raw,
                    "connection_id": display_connection_id,
                    "base_connection_id": connection_id_raw,
                    "connection_name": pick(row, "connection_name"),
                    "direction": normalize_direction(pick(row, "direction")),
                    "connection_attribute": pick(row, "connection_attribute"),
                    "link_family_id": link_meta.get("link_family_id", ""),
                    "link_instance_id": link_meta.get("link_instance_id", ""),
                    "link_member_sheets": link_meta.get("link_member_sheets", []),
                    "user_link_info": link_meta.get("user_link_info", ""),
                    "device_role_info": link_meta.get("device_role_info", ""),
                    "link_contexts": link_meta.get("link_contexts", []),
                    "raw_source_port": source_port,
                    "raw_target_port": target_port_raw,
                }
                out.append(normalized)

    write_jsonl(output_path, out)
    return out

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--connections", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()
    rows = normalize_connections(args.connections, args.output)
    print(f"[OK] normalized connections: {len(rows)} -> {args.output}")

if __name__ == "__main__":
    main()
