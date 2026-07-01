#!/usr/bin/env python3
"""SIP-backed, deterministic gates for pipeline-schematic Phase A."""

from __future__ import annotations

import argparse
import json
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook
from sip import StateIntegrityProtocol


ANCHORS = {
    "A4_INPUT_READY": (
        "A4 input block diagram aggregation and pin catalog are complete "
        "readable and cover every device part before semantic mapping"
    ),
    "A4_REAL_SUBAGENT_ANALYSIS": (
        "A4 semantic mapping decisions are produced by isolated real "
        "subagents and never substituted by stage all or analysis scripts"
    ),
    "A4_OUTPUT_CONTRACT": (
        "A4 final signal interface preserves workbook pagination sheet order "
        "information sheets and renders every connection sheet with 13 columns"
    ),
    "PHASE_A_A5_COMPLETED": (
        "Phase A completes only after A5 net consistency check fix and recheck "
        "produce a readable checked workbook with zero proposed changes and errors"
    ),
}

INFO_SHEETS = {"BLOCK_INFO", "LINK_INFO", "链路信息", "说明", "README", "INDEX", "目录"}
PART_HEADERS = {
    "器件信息",
    "器件编码",
    "器件料号",
    "part_number",
    "part_id",
    "device_part_id",
    "device_code",
    "source_part_id",
}


def fail(code: str, message: str, evidence: str | None = None) -> dict[str, str]:
    item = {"code": code, "message": message}
    if evidence:
        item["evidence"] = evidence
    return item


def existing_file(path: Path | None, code: str, errors: list[dict[str, str]]) -> bool:
    if path is None or not path.is_file() or path.stat().st_size == 0:
        errors.append(fail(code, "required non-empty file is missing", str(path)))
        return False
    return True


def workbook(path: Path, errors: list[dict[str, str]]):
    try:
        return load_workbook(path, read_only=True, data_only=False)
    except Exception as exc:
        errors.append(fail("WORKBOOK_UNREADABLE", str(exc), str(path)))
        return None


def normalized(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def block_parts(ws) -> set[str]:
    rows = ws.iter_rows(values_only=True)
    try:
        headers = [normalized(v) for v in next(rows)]
    except StopIteration:
        return set()
    indexes = [i for i, header in enumerate(headers) if header in PART_HEADERS]
    if not indexes:
        return set()
    parts: set[str] = set()
    for row in rows:
        for index in indexes:
            if index < len(row) and normalized(row[index]):
                parts.add(normalized(row[index]))
    return parts


def resolve_pin_key(catalog: dict[str, Any], part: str) -> str | None:
    if part in catalog:
        return part
    upper = part.upper()
    for key in catalog:
        if normalized(key).upper() == upper:
            return key
    return None


def gate_input(args, errors: list[dict[str, str]], evidence: list[dict[str, Any]]) -> None:
    if not existing_file(args.block_info, "A4_INPUT_MISSING", errors):
        return
    if not existing_file(args.pin_info, "A4_INPUT_MISSING", errors):
        return
    wb = workbook(args.block_info, errors)
    if wb is None:
        return
    try:
        evidence.append({"path": str(args.block_info.resolve()), "sheets": wb.sheetnames})
        if "BLOCK_INFO" not in wb.sheetnames:
            errors.append(fail("A4_BLOCK_INFO_MISSING", "BLOCK_INFO sheet is missing"))
            return
        connection_sheets = [name for name in wb.sheetnames if name not in INFO_SHEETS]
        if not connection_sheets:
            errors.append(fail("A4_CONNECTION_SHEET_MISSING", "no connection sheet found"))
        parts = block_parts(wb["BLOCK_INFO"])
        if not parts:
            errors.append(
                fail("A4_DEVICE_PARTS_MISSING", "BLOCK_INFO has no recognized non-empty part column")
            )
            return
    finally:
        wb.close()
    try:
        catalog = json.loads(args.pin_info.read_text(encoding="utf-8-sig"))
    except Exception as exc:
        errors.append(fail("A4_PIN_JSON_INVALID", str(exc), str(args.pin_info)))
        return
    if not isinstance(catalog, dict):
        errors.append(fail("A4_PIN_JSON_INVALID", "pin_info root must be an object"))
        return
    missing = []
    for part in sorted(parts):
        key = resolve_pin_key(catalog, part)
        if key is None or not isinstance(catalog[key], list) or not catalog[key]:
            missing.append(part)
    evidence.append(
        {
            "path": str(args.pin_info.resolve()),
            "device_parts": sorted(parts),
            "covered_parts": len(parts) - len(missing),
        }
    )
    if missing:
        errors.append(
            fail(
                "A4_PIN_COVERAGE_INCOMPLETE",
                f"missing or empty pin lists: {', '.join(missing)}",
                str(args.pin_info),
            )
        )


def walk_objects(value: Any) -> Iterable[dict[str, Any]]:
    if isinstance(value, dict):
        yield value
        for child in value.values():
            yield from walk_objects(child)
    elif isinstance(value, list):
        for child in value:
            yield from walk_objects(child)


def read_json(path: Path, code: str, errors: list[dict[str, str]]) -> Any:
    try:
        return json.loads(path.read_text(encoding="utf-8-sig"))
    except Exception as exc:
        errors.append(fail(code, str(exc), str(path)))
        return None


def task_contracts(plan: Any) -> list[tuple[Path, set[str]]]:
    contracts: list[tuple[Path, set[str]]] = []
    for item in walk_objects(plan):
        contract = item.get("output_contract")
        if not isinstance(contract, dict) or not contract.get("output_file"):
            continue
        raw_ids = (
            contract.get("required_line_ids")
            or contract.get("line_ids")
            or item.get("line_ids")
            or item.get("required_line_ids")
            or []
        )
        contracts.append((Path(str(contract["output_file"])), {normalized(v) for v in raw_ids}))
    return contracts


def gate_subagents(args, errors: list[dict[str, str]], evidence: list[dict[str, Any]]) -> None:
    required = [
        (args.session_plan, "A4_SESSION_PLAN_MISSING"),
        (args.task_plan_json, "A4_TASK_PLAN_MISSING"),
        (args.task_plan_md, "A4_TASK_PLAN_MISSING"),
    ]
    if not all(existing_file(path, code, errors) for path, code in required):
        return
    text = args.task_plan_md.read_text(encoding="utf-8-sig").lower()
    if "execution_mode=subagent" not in text and "execution_mode:subagent" not in text:
        errors.append(
            fail(
                "A4_SUBAGENT_NOT_USED",
                "TASK_PLAN.md does not record execution_mode=subagent",
                str(args.task_plan_md),
            )
        )
    forbidden = ("stage=all", "--stage all", "execution_mode=script")
    found = [token for token in forbidden if token in text]
    if found:
        errors.append(
            fail(
                "A4_SCRIPT_SUBSTITUTED_ANALYSIS",
                f"forbidden execution markers: {', '.join(found)}",
                str(args.task_plan_md),
            )
        )
    plan = read_json(args.task_plan_json, "A4_TASK_PLAN_INVALID", errors)
    sessions = read_json(args.session_plan, "A4_SESSION_PLAN_INVALID", errors)
    if plan is None or sessions is None:
        return
    contracts = task_contracts(plan)
    if not contracts:
        errors.append(fail("A4_TASK_CONTRACT_MISSING", "no output_contract.output_file found"))
        return
    checked = []
    for output_file, required_ids in contracts:
        path = output_file if output_file.is_absolute() else args.task_plan_json.parent / output_file
        if not existing_file(path, "A4_SUBAGENT_OUTPUT_MISSING", errors):
            continue
        seen: list[str] = []
        try:
            for number, line in enumerate(path.read_text(encoding="utf-8-sig").splitlines(), 1):
                if not line.strip():
                    continue
                obj = json.loads(line)
                line_id = normalized(obj.get("line_id"))
                if not line_id:
                    raise ValueError(f"line {number}: line_id is empty")
                seen.append(line_id)
        except Exception as exc:
            errors.append(fail("A4_SUBAGENT_OUTPUT_INVALID", str(exc), str(path)))
            continue
        if len(seen) != len(set(seen)):
            errors.append(fail("A4_SUBAGENT_LINE_ID_DUPLICATE", "duplicate line_id", str(path)))
        if required_ids and set(seen) != required_ids:
            errors.append(
                fail(
                    "A4_SUBAGENT_LINE_ID_COVERAGE",
                    f"expected {sorted(required_ids)}, observed {sorted(set(seen))}",
                    str(path),
                )
            )
        checked.append(str(path.resolve()))
    evidence.append(
        {
            "session_plan": str(args.session_plan.resolve()),
            "task_contract_count": len(contracts),
            "checked_outputs": checked,
        }
    )


def sheet_values(ws) -> list[list[Any]]:
    return [list(row) for row in ws.iter_rows(values_only=True)]


def gate_output(args, errors: list[dict[str, str]], evidence: list[dict[str, Any]]) -> None:
    required = [
        (args.block_info, "A4_INPUT_MISSING"),
        (args.signal_interface, "A4_OUTPUT_MISSING"),
        (args.validation_report, "A4_VALIDATION_REPORT_MISSING"),
        (args.subagent_check, "A4_SUBAGENT_CHECK_MISSING"),
    ]
    if not all(existing_file(path, code, errors) for path, code in required):
        return
    source = workbook(args.block_info, errors)
    output = workbook(args.signal_interface, errors)
    if source is None or output is None:
        return
    try:
        if source.sheetnames != output.sheetnames:
            errors.append(
                fail(
                    "A4_OUTPUT_PAGINATION_INVALID",
                    f"expected sheets {source.sheetnames}, observed {output.sheetnames}",
                )
            )
        for name in output.sheetnames:
            if name in INFO_SHEETS and name in source.sheetnames:
                if sheet_values(source[name]) != sheet_values(output[name]):
                    errors.append(
                        fail("A4_INFORMATION_SHEET_CHANGED", f"{name} was not preserved")
                    )
            elif name not in INFO_SHEETS:
                ws = output[name]
                if ws.max_column != 13:
                    errors.append(
                        fail(
                            "A4_OUTPUT_COLUMN_COUNT_INVALID",
                            f"{name} has {ws.max_column} columns, expected 13",
                        )
                    )
                if ws.max_row < 2:
                    errors.append(fail("A4_OUTPUT_CONNECTIONS_EMPTY", f"{name} has no data rows"))
        evidence.append(
            {
                "input": str(args.block_info.resolve()),
                "output": str(args.signal_interface.resolve()),
                "sheets": output.sheetnames,
            }
        )
    finally:
        source.close()
        output.close()
    for path, code in (
        (args.validation_report, "A4_VALIDATION_ERROR"),
        (args.subagent_check, "A4_SUBAGENT_CHECK_FAILED"),
    ):
        report = read_json(path, code, errors)
        if report is not None and report_has_error(report):
            errors.append(fail(code, "report contains ERROR/failed status", str(path)))


def report_has_error(value: Any) -> bool:
    for obj in walk_objects(value):
        for key, child in obj.items():
            key_lower = str(key).lower()
            if key_lower in {"status", "overall_status", "severity", "level"}:
                if normalized(child).lower() in {"error", "failed", "fail"}:
                    return True
            if key_lower in {"errors", "error_count"}:
                if isinstance(child, list) and child:
                    return True
                if isinstance(child, (int, float)) and child > 0:
                    return True
    return False


def find_numbers(value: Any, wanted: set[str]) -> list[float]:
    found: list[float] = []
    if isinstance(value, dict):
        for key, child in value.items():
            if str(key).lower() in wanted and isinstance(child, (int, float)):
                found.append(float(child))
            found.extend(find_numbers(child, wanted))
    elif isinstance(value, list):
        for child in value:
            found.extend(find_numbers(child, wanted))
    return found


def gate_a5(args, errors: list[dict[str, str]], evidence: list[dict[str, Any]]) -> None:
    required = [
        (args.a5_excel, "PHASE_A_A5_OUTPUT_MISSING"),
        (args.a5_report, "PHASE_A_A5_REPORT_MISSING"),
        (args.task_plan_md, "PHASE_A_TASK_PLAN_MISSING"),
    ]
    if not all(existing_file(path, code, errors) for path, code in required):
        return
    checked = workbook(args.a5_excel, errors)
    if checked is not None:
        checked.close()
    plan = args.task_plan_md.read_text(encoding="utf-8-sig").lower()
    if "a5" not in plan or "completed" not in plan or "diagram-signal-net-consistency-checker" not in plan:
        errors.append(
            fail(
                "PHASE_A_A5_NOT_RUN",
                "TASK_PLAN.md lacks A5 completed checker evidence",
                str(args.task_plan_md),
            )
        )
    report = read_json(args.a5_report, "PHASE_A_A5_REPORT_INVALID", errors)
    if report is None:
        return
    counts = find_numbers(report, {"proposed_change_count"})
    if not counts:
        errors.append(
            fail("PHASE_A_A5_REPORT_INVALID", "proposed_change_count is absent", str(args.a5_report))
        )
    elif counts[-1] != 0:
        errors.append(
            fail(
                "PHASE_A_A5_INCOMPLETE",
                f"final proposed_change_count is {counts[-1]}, expected 0",
                str(args.a5_report),
            )
        )
    if report_has_error(report):
        errors.append(fail("PHASE_A_A5_INCOMPLETE", "A5 report contains ERROR", str(args.a5_report)))
    evidence.append(
        {
            "checked_excel": str(args.a5_excel.resolve()),
            "report": str(args.a5_report.resolve()),
            "proposed_change_counts": counts,
        }
    )


def append_task_plan(path: Path | None, result: dict[str, Any]) -> None:
    if path is None:
        return
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("a", encoding="utf-8") as handle:
        handle.write("\n[STATE_INTEGRITY_ANCHOR]\n")
        handle.write(json.dumps(result, ensure_ascii=False, sort_keys=True) + "\n")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--anchor-id", required=True, choices=sorted(ANCHORS))
    parser.add_argument("--observation", required=True)
    parser.add_argument("--threshold", type=float, default=0.35)
    parser.add_argument("--result-json", type=Path)
    parser.add_argument("--task-plan-md", type=Path)
    parser.add_argument("--block-info", type=Path)
    parser.add_argument("--pin-info", type=Path)
    parser.add_argument("--session-plan", type=Path)
    parser.add_argument("--task-plan-json", type=Path)
    parser.add_argument("--signal-interface", type=Path)
    parser.add_argument("--validation-report", type=Path)
    parser.add_argument("--subagent-check", type=Path)
    parser.add_argument("--a5-excel", type=Path)
    parser.add_argument("--a5-report", type=Path)
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    errors: list[dict[str, str]] = []
    evidence: list[dict[str, Any]] = []
    dispatch = {
        "A4_INPUT_READY": gate_input,
        "A4_REAL_SUBAGENT_ANALYSIS": gate_subagents,
        "A4_OUTPUT_CONTRACT": gate_output,
        "PHASE_A_A5_COMPLETED": gate_a5,
    }
    try:
        sip = StateIntegrityProtocol(threshold=args.threshold)
        sip.anchor(ANCHORS[args.anchor_id])
        observation = sip.observe(args.observation)
        sip_result = {
            "available": True,
            "version": __import__("sip").__version__,
            "drift": observation.drift,
            "threshold": observation.threshold,
            "aligned": observation.is_aligned,
        }
        if not observation.is_aligned:
            errors.append(
                fail(
                    "SIP_SEMANTIC_DRIFT",
                    f"drift {observation.drift:.6f} exceeds {observation.threshold:.6f}",
                )
            )
    except Exception as exc:
        sip_result = {"available": False, "error": str(exc)}
        errors.append(fail("SIP_RUNTIME_ERROR", str(exc)))
    dispatch[args.anchor_id](args, errors, evidence)
    result = {
        "anchor_id": args.anchor_id,
        "status": "PASS" if not errors else "FAIL",
        "timestamp": datetime.now(timezone.utc).isoformat(),
        "sip": sip_result,
        "evidence": evidence,
        "failure_codes": [item["code"] for item in errors],
        "errors": errors,
        "next_step": "CONTINUE" if not errors else "HALT",
    }
    rendered = json.dumps(result, ensure_ascii=False, indent=2)
    print(rendered)
    if args.result_json:
        args.result_json.parent.mkdir(parents=True, exist_ok=True)
        args.result_json.write_text(rendered + "\n", encoding="utf-8")
    append_task_plan(args.task_plan_md, result)
    return 0 if not errors else 2


if __name__ == "__main__":
    sys.exit(main())
