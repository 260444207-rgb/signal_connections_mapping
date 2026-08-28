from __future__ import annotations

from copy import copy
import json
import re
import sys
import tempfile
import unittest
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook, load_workbook


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from apply_naming import apply_naming
from prepare_naming import load_model_rule_bundle, load_model_rules, prepare_naming


HEADERS = [
    "源Block标识", "源Block名称", "源Port", "目的Block标识", "目的Block名称",
    "目的Port", "连线ID", "连线名称", "连线方向", "原理图Pin脚", "分析说明",
    "映射置信度", "网络命名",
]


def make_workbook(path: Path, rows: list[list[str]]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "DEV"
    sheet.append(HEADERS)
    for row in rows:
        sheet.append(row)
    workbook.save(path)
    workbook.close()


def add_ignorable_namespace(path: Path) -> None:
    temporary = path.with_suffix(".namespaced.xlsx")
    with ZipFile(path, "r") as source, ZipFile(temporary, "w", ZIP_DEFLATED) as target:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                data = data.replace(
                    b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"',
                    b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
                    b'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
                    b'xmlns:x14ac="http://schemas.microsoft.com/office/spreadsheetml/2009/9/ac" '
                    b'mc:Ignorable="x14ac"',
                    1,
                )
            target.writestr(copy(info), data)
    temporary.replace(path)


def worksheet_root(path: Path) -> bytes:
    with ZipFile(path) as archive:
        data = archive.read("xl/worksheets/sheet1.xml")
    return re.search(rb"<worksheet\b[^>]*>", data).group(0)


class NetworkNamingTests(unittest.TestCase):
    def test_nonempty_connection_name_is_used_verbatim_without_model(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "input.xlsx"
            output = root / "output.xlsx"
            task_dir = root / "tasks"
            make_workbook(source, [[
                "B1", "SRC", "OUT", "B2", "DST", "IN", "C1", "valid_name_1",
                "INPUT", "PIN_A", "pin 分析依据", "High", "",
            ]])

            report = prepare_naming(source, task_dir)
            self.assertEqual("PASS", report["status"])
            self.assertEqual(1, report["automatic_group_count"])
            self.assertEqual(0, report["model_group_count"])
            self.assertEqual("", (task_dir / "naming_groups.jsonl").read_text(encoding="utf-8"))
            automatic = json.loads((task_dir / "automatic_decisions.jsonl").read_text(encoding="utf-8"))
            self.assertEqual("valid_name_1", automatic["net_name"])

            result = apply_naming(source, task_dir, output)
            self.assertEqual("PASS", result["status"])
            workbook = load_workbook(output)
            self.assertEqual("valid_name_1", workbook["DEV"]["M2"].value)
            self.assertIn("pin 分析依据；网络命名依据：连线名称非空且为合法网络名", workbook["DEV"]["K2"].value)
            workbook.close()

    def test_invalid_connection_names_create_two_end_model_task_with_type_rules(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "input.xlsx"
            output = root / "output.xlsx"
            task_dir = root / "tasks"
            make_workbook(source, [
                ["B1", "SRC", "CTRL", "B2", "DST", "EN", "C2", "LINE_175", "OUTPUT", "PIN_CTRL", "该 pin 用于目标使能控制", "High", ""],
                ["B2", "DST", "EN", "B1", "SRC", "CTRL", "C2", "BAD NAME", "INPUT", "PIN_EN", "目标端实际使能 pin", "High", ""],
            ])

            report = prepare_naming(source, task_dir)
            self.assertEqual(0, report["automatic_group_count"])
            self.assertEqual(1, report["model_group_count"])
            task = json.loads((task_dir / "naming_groups.jsonl").read_text(encoding="utf-8"))
            self.assertIn("该 pin 用于目标使能控制", task["groups"][0]["analysis_notes"])
            self.assertEqual(load_model_rules(), task["rules"])
            self.assertEqual(["DIGITAL", "RF", "POWER", "GROUND"], task["signal_types"])
            bundle = load_model_rule_bundle()
            self.assertEqual(bundle["classification"], task["classification_rules"])
            self.assertEqual(bundle["by_signal_type"], task["type_naming_rules"])
            self.assertEqual({"DIGITAL", "RF", "POWER", "GROUND"}, set(task["type_naming_rules"]))
            details = task["groups"][0]["connection_details"][0]
            self.assertEqual(["LINE_175", "BAD NAME"], details["connection_names"])
            self.assertEqual({"PIN_CTRL", "PIN_EN"}, {endpoint["pin"] for endpoint in details["endpoints"]})
            (task_dir / "naming_decisions.jsonl").write_text(
                json.dumps({"id": "G0001", "signal_type": "DIGITAL", "net_name": "SRC_DST_EN", "basis": "结合两端 pin 判断为使能信号。"}, ensure_ascii=False) + "\n",
                encoding="utf-8",
            )
            result = apply_naming(source, task_dir, output)
            self.assertEqual("PASS", result["status"])
            workbook = load_workbook(output)
            self.assertEqual("SRC_DST_EN", workbook["DEV"]["M2"].value)
            workbook.close()

    def test_conflicting_connection_names_in_one_group_stop_prepare(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "input.xlsx"
            task_dir = root / "tasks"
            make_workbook(source, [
                ["B1", "SRC", "OUT", "B2", "DST", "IN", "C1", "NAME_A", "OUTPUT", "PIN_A", "A", "High", ""],
                ["B1", "SRC", "OUT", "B3", "DST2", "IN", "C1", "NAME_B", "OUTPUT", "PIN_A", "B", "High", ""],
            ])
            report = prepare_naming(source, task_dir)
            self.assertEqual("ERROR", report["status"])
            self.assertTrue(any("多个不同的合法连线名称" in error["message"] for error in report["errors"]))

    def test_direction_does_not_resolve_connection_name_conflict(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "input.xlsx"
            task_dir = root / "tasks"
            make_workbook(source, [
                ["B1", "SRC", "OUT", "B2", "DST", "IN", "C1", "OUTPUT_NAME", "OUTPUT", "PIN_A", "A", "High", ""],
                ["B2", "DST", "IN", "B1", "SRC", "OUT", "C1", "INPUT_NAME", "INPUT", "PIN_B", "B", "High", ""],
            ])
            report = prepare_naming(source, task_dir)
            self.assertEqual("ERROR", report["status"])
            self.assertEqual(0, report["automatic_group_count"])
            self.assertEqual(0, report["model_group_count"])
            self.assertTrue(any("OUTPUT_NAME" in error["message"] and "INPUT_NAME" in error["message"] for error in report["errors"]))

    def test_model_decision_requires_signal_type(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "input.xlsx"
            output = root / "output.xlsx"
            task_dir = root / "tasks"
            make_workbook(source, [[
                "B1", "SRC", "OUT", "B2", "DST", "IN", "C1", "LINE_1",
                "OUTPUT", "PIN_A", "analysis", "High", "",
            ]])
            prepare_naming(source, task_dir)
            (task_dir / "naming_decisions.jsonl").write_text(
                json.dumps({"id": "G0001", "net_name": "SRC_DST_SIG", "basis": "测试。"}, ensure_ascii=False) + "\n",
                encoding="utf-8",
            )
            result = apply_naming(source, task_dir, output)
            self.assertEqual("ERROR", result["status"])
            self.assertTrue(any("signal_type" in error["message"] for error in result["errors"]))

    def test_output_suffix_grouping_still_requires_output_direction(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "input.xlsx"
            task_dir = root / "tasks"
            make_workbook(source, [
                ["B1", "SRC", "OUT", "B2", "A", "IN", "LEFT_1", "", "INPUT", "PIN_X", "A", "High", ""],
                ["B1", "SRC", "OUT", "B3", "B", "IN", "RIGHT_1", "", "INPUT", "PIN_X", "B", "High", ""],
            ])
            report = prepare_naming(source, task_dir)
            self.assertEqual(2, report["group_count"])
            self.assertEqual(2, report["model_group_count"])

    def test_cell_patch_preserves_extension_namespace_and_opens(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "input.xlsx"
            output = root / "output.xlsx"
            task_dir = root / "tasks"
            make_workbook(source, [[
                "B1", "SRC", "OUT", "B2", "DST", "IN", "C1", "DIRECT_NAME",
                "OUTPUT", "PIN_A", "analysis", "High", "",
            ]])
            add_ignorable_namespace(source)
            original_root = worksheet_root(source)
            prepare_naming(source, task_dir)
            result = apply_naming(source, task_dir, output)
            self.assertEqual("PASS", result["status"])
            self.assertEqual(original_root, worksheet_root(output))
            with ZipFile(output) as archive:
                self.assertIsNone(archive.testzip())
            workbook = load_workbook(output)
            self.assertEqual("DIRECT_NAME", workbook["DEV"]["M2"].value)
            workbook.close()


if __name__ == "__main__":
    unittest.main()
